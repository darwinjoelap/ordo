"""Importación de clientes desde Excel."""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.clientes import importacion as imp
from apps.clientes.models import Cliente, ImportacionClientes

from .test_fase4 import Base


def excel(filas, titulos=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Clientes'
    ws.append(titulos or imp.CLAVES)
    for f in filas:
        ws.append(f)
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


class ImportarClientesTests(Base):
    """Ya existen: Clínica Sol (J-1, vendedor) y Lab Luna (J-2)."""

    def validar(self, filas, **k):
        with self.empresa_ctx():
            return imp.validar(excel(filas, **k), self.empresa)

    def test_plantilla_se_puede_reimportar(self):
        wb = load_workbook(BytesIO(imp.generar_plantilla([('v', 'Vendedor')])))
        self.assertEqual(wb.sheetnames, ['Clientes', 'Instrucciones', 'Vendedores'])
        self.assertEqual([c.value for c in wb['Clientes'][1]], imp.CLAVES)

    def test_crea_limpia_y_asigna_vendedor(self):
        r = self.validar([['AGRO\xa0 DON  EULOGIO, C.A', 'j-999', 'Ana', '0414-1', 'A@B.COM', 'Quíbor', 'nota', 'V@ALFA.COM', 'no'],
                          ['Solo Nombre', None, None, None, None, None, None, None, None]])
        self.assertTrue(r.valido, r.errores)
        self.assertEqual(r.resumen()['nuevos'], 2)
        self.assertEqual(r.resumen()['sin_rif'], 1)
        with self.empresa_ctx():
            self.assertEqual(imp.aplicar(r), {'creados': 2, 'actualizados': 0})
            c = Cliente.objects.get(rif='J-999')
            self.assertEqual((c.nombre, c.email, c.vendedor, c.activo), ('AGRO DON EULOGIO, C.A', 'a@b.com', self.vendedor, False))
            self.assertTrue(Cliente.objects.get(nombre='Solo Nombre').activo)

    def test_actualiza_por_rif_o_por_nombre_sin_borrar(self):
        r = self.validar([['Clínica Sol, C.A.', 'J1', None, '0251-1', None, None, None, None, None],
                          ['LAB LUNA', None, 'Luis', None, None, None, None, None, None]])
        self.assertTrue(r.valido, r.errores)
        self.assertEqual(r.resumen()['actualizados'], 2)
        with self.empresa_ctx():
            imp.aplicar(r)
            self.assertEqual(Cliente.objects.count(), 2)
            sol, luna = Cliente.objects.get(pk=self.cliente.pk), Cliente.objects.get(pk=self.cliente2.pk)
        self.assertEqual((sol.nombre, sol.rif, sol.telefono, sol.vendedor), ('Clínica Sol, C.A.', 'J1', '0251-1', self.vendedor))
        self.assertEqual((luna.rif, luna.contacto), ('J-2', 'Luis'))

    def test_errores(self):
        r = self.validar([[None, 'J-5', None, None, None, None, None, None, None],
                          ['A', 'J-7', None, None, 'malo', None, None, 'nadie', 'quizás'],
                          ['B', 'j7', None, None, None, None, None, None, None],
                          ['Repetido', None, None, None, None, None, None, None, None],
                          ['REPETIDO.', None, None, None, None, None, None, None, None]])
        columnas = sorted((f, c) for f, c, _ in r.errores)
        self.assertEqual(columnas, [(2, 'nombre'), (3, 'activo'), (3, 'email'), (3, 'vendedor'), (4, 'rif'), (6, 'nombre')])
        with self.assertRaises(imp.ErrorArchivo):
            self.validar([['x']], titulos=['razon'])

    def test_no_ve_clientes_de_otra_empresa(self):
        from apps.core.tenancy import usando_empresa
        with usando_empresa(self.otra):
            Cliente.objects.create(nombre='Ajeno', rif='J-88')
        r = self.validar([['Ajeno', 'J-88', None, None, None, None, None, None, None]])
        self.assertEqual(r.resumen()['nuevos'], 1)

    def test_pantallas(self):
        self.client.force_login(self.vendedor)
        self.assertIn(self.client.get(reverse('clientes:importar')).status_code, (302, 403))
        self.assertNotContains(self.client.get(reverse('clientes:lista')), reverse('clientes:importar'))
        self.client.force_login(self.dueno)
        self.assertContains(self.client.get(reverse('clientes:lista')), reverse('clientes:importar'))
        self.assertContains(self.client.get(reverse('clientes:importar')), 'Descargar plantilla')
        self.assertTrue(self.client.get(reverse('clientes:plantilla')).content.startswith(b'PK'))
        self.assertTrue(self.client.get(reverse('clientes:exportar')).content.startswith(b'PK'))
        sube = lambda filas: self.client.post(reverse('clientes:importar'), {
            'archivo': SimpleUploadedFile('c.xlsx', excel(filas))}, follow=True)
        r = sube([[None, 'J-5', None, None, None, None, None, None, None]])
        self.assertContains(r, 'No se importó nada')
        r = sube([['Nuevo Cliente', 'J-50', None, None, None, None, None, None, None]])
        self.assertContains(r, 'Importar ahora')
        i = ImportacionClientes.todos.get(estado='VALIDADA')
        r = self.client.post(reverse('clientes:importar_confirmar', args=[i.pk]), follow=True)
        self.assertContains(r, '1 clientes nuevos')
        self.assertTrue(Cliente.todos.filter(rif='J-50', empresa=self.empresa).exists())
        self.assertEqual(self.client.post(reverse('clientes:importar_confirmar', args=[i.pk])).status_code, 302)
        self.assertEqual(Cliente.todos.filter(rif='J-50').count(), 1)
