"""
Pruebas de la Fase 3B: importación Excel, compras, panel de pedido y lista de precios.
"""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.compras import servicios as compras
from apps.compras import sugerencias
from apps.compras.models import ItemOrdenCompra, OrdenCompra
from apps.core.tenancy import usando_empresa
from apps.empresas.models import Empresa, Membresia, Rol
from apps.inventario import importacion, servicios
from apps.inventario.models import (Categoria, ImportacionProductos, Lote, Marca, MovimientoInventario, Producto,
                                    Unidad)
from apps.proveedores.models import Proveedor

HOY = timezone.localdate()


def excel(filas, encabezado=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Productos'
    ws.append(encabezado or importacion.CLAVES)
    for f in filas:
        ws.append([f.get(c) for c in (encabezado or importacion.CLAVES)])
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


class Base(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre='Alfa')
        self.otra = Empresa.objects.create(nombre='Beta')
        U = get_user_model()
        self.dueno = U.objects.create_user('d@alfa.com', 'clave-segura-123')
        self.vendedor = U.objects.create_user('v@alfa.com', 'clave-segura-123')
        Membresia.objects.create(usuario=self.dueno, empresa=self.empresa, rol=Rol.DUENO)
        Membresia.objects.create(usuario=self.vendedor, empresa=self.empresa, rol=Rol.VENDEDOR)
        self.client.force_login(self.dueno)

    def en_empresa(self):
        return usando_empresa(self.empresa)


# ── Importación ───────────────────────────────────────────────────────────────

class PlantillaTests(Base):
    def test_descarga_plantilla_con_hojas_y_columnas(self):
        r = self.client.get(reverse('inventario:plantilla'))
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Productos', 'Instrucciones', 'Unidades'])
        self.assertEqual([c.value for c in wb['Productos'][1]], importacion.CLAVES)
        unidades = [c.value for c in wb['Unidades']['A']][1:]
        self.assertIn('KG', unidades)       # unidades reales de la empresa

    def test_la_propia_plantilla_de_ejemplo_es_valida(self):
        with self.en_empresa():
            r = importacion.validar(importacion.generar_plantilla([('UN', 'Unidad'), ('KIT', 'Kit'), ('CAJA', 'Caja')]))
        self.assertTrue(r.valido, r.errores)
        self.assertEqual(r.resumen()['productos_nuevos'], 2)
        self.assertEqual(r.resumen()['lotes'], 3)

    def test_vendedor_no_puede_importar(self):
        self.client.force_login(self.vendedor)
        self.assertEqual(self.client.get(reverse('inventario:importar')).status_code, 403)


class ValidacionTests(Base):
    def validar(self, filas, **kw):
        with self.en_empresa():
            return importacion.validar(excel(filas, **kw))

    def test_errores_por_fila_y_columna(self):
        r = self.validar([
            {'codigo': 'A1', 'nombre': '', 'categoria': 'X'},
            {'codigo': 'A2', 'nombre': 'Ok', 'categoria': 'X', 'unidad': 'LITROS'},
            {'codigo': 'A3', 'nombre': 'Ok', 'categoria': 'X', 'costo_usd': 'diez'},
            {'codigo': 'A4', 'nombre': 'Ok', 'categoria': 'X', 'maneja_lote': 'quizás'},
            {'codigo': 'A5', 'nombre': 'Ok', 'categoria': 'X', 'maneja_lote': 'SI', 'existencia_inicial': 5},
            {'codigo': 'A6', 'nombre': 'Ok', 'categoria': 'X', 'existencia_inicial': 2, 'fecha_vencimiento': '31/13/2027'},
        ])
        self.assertFalse(r.valido)
        columnas = {(f, c) for f, c, _ in r.errores}
        self.assertEqual(columnas, {(2, 'nombre'), (3, 'unidad'), (4, 'costo_usd'), (5, 'maneja_lote'),
                                    (6, 'numero_lote'), (7, 'fecha_vencimiento')})

    def test_formatos_flexibles(self):
        r = self.validar([{'codigo': 'b-1', 'nombre': 'B', 'categoria': 'X', 'costo_usd': '1.234,50',
                           'precio_venta_usd': '12,5', 'maneja_lote': 'sí', 'activo': 'no',
                           'existencia_inicial': '3', 'numero_lote': 'L1', 'fecha_vencimiento': '05/02/2027'}])
        self.assertTrue(r.valido, r.errores)
        p = r.productos['B-1']
        self.assertEqual(p.datos['costo_usd'], Decimal('1234.50'))
        self.assertEqual(p.datos['precio_venta_usd'], Decimal('12.50'))
        self.assertEqual((p.datos['maneja_lote'], p.datos['activo']), (True, False))
        self.assertEqual(p.lotes, [(2, 3, 'L1', date(2027, 2, 5))])

    def test_varios_lotes_del_mismo_producto(self):
        r = self.validar([
            {'codigo': 'G1', 'nombre': 'Glucosa', 'categoria': 'R', 'maneja_lote': 'SI',
             'existencia_inicial': 5, 'numero_lote': 'L1'},
            {'codigo': 'G1', 'existencia_inicial': 7, 'numero_lote': 'L2'},
            {'codigo': 'G1', 'existencia_inicial': 1, 'numero_lote': 'L2'},
        ])
        self.assertEqual([(f, c) for f, c, _ in r.errores], [(4, 'numero_lote')])   # lote repetido

    def test_faltan_columnas_obligatorias(self):
        with self.assertRaises(importacion.ErrorArchivo):
            self.validar([{'codigo': 'A'}], encabezado=['codigo', 'precio'])

    def test_archivo_que_no_es_excel(self):
        with self.en_empresa(), self.assertRaises(importacion.ErrorArchivo):
            importacion.validar(b'esto no es un excel')


class AplicarTests(Base):
    def test_flujo_completo_desde_la_vista(self):
        contenido = excel([
            {'codigo': 'G1', 'nombre': 'Glucosa', 'categoria': 'Reactivos', 'subcategoria': 'Química',
             'marca': 'Wiener', 'unidad': 'kit', 'proveedor': 'Distri Lara', 'costo_usd': 40, 'precio_venta_usd': 80,
             'maneja_lote': 'SI', 'maneja_vencimiento': 'SI', 'existencia_inicial': 5, 'numero_lote': 'L1',
             'fecha_vencimiento': date(2027, 1, 31)},
            {'codigo': 'G1', 'existencia_inicial': 3, 'numero_lote': 'L2', 'fecha_vencimiento': date(2027, 6, 30)},
            {'codigo': 'T1', 'nombre': 'Tubos', 'categoria': 'Consumibles'},
        ])
        r = self.client.post(reverse('inventario:importar'),
                             {'archivo': SimpleUploadedFile('productos.xlsx', contenido)})
        imp = ImportacionProductos.todos.get()
        self.assertRedirects(r, reverse('inventario:importar_revision', args=[imp.pk]))
        self.assertEqual(imp.estado, 'VALIDADA')
        self.assertFalse(Producto.todos.exists())             # validar no guarda nada
        r = self.client.get(reverse('inventario:importar_revision', args=[imp.pk]))
        self.assertContains(r, 'Importar ahora')
        self.client.post(reverse('inventario:importar_confirmar', args=[imp.pk]))
        with self.en_empresa():
            g = Producto.objects.con_stock().get(codigo='G1')
            self.assertEqual((g.anot_stock_total, g.subcategoria.nombre, g.marca.nombre, g.unidad.abreviatura,
                              g.proveedor_habitual.nombre), (8, 'Química', 'Wiener', 'KIT', 'Distri Lara'))
            self.assertEqual(Lote.objects.filter(producto=g).count(), 2)
            self.assertEqual(MovimientoInventario.objects.count(), 2)
        imp.refresh_from_db()
        self.assertEqual((imp.estado, bytes(imp.contenido)), ('APLICADA', b''))
        # No se puede aplicar dos veces
        self.client.post(reverse('inventario:importar_confirmar', args=[imp.pk]))
        self.assertEqual(Producto.todos.count(), 2)

    def test_con_errores_no_ofrece_importar(self):
        r = self.client.post(reverse('inventario:importar'),
                             {'archivo': SimpleUploadedFile('p.xlsx', excel([{'codigo': 'X'}]))}, follow=True)
        self.assertContains(r, 'No se importó nada')
        self.assertNotContains(r, 'Importar ahora')

    def test_rechaza_csv(self):
        r = self.client.post(reverse('inventario:importar'),
                             {'archivo': SimpleUploadedFile('p.csv', b'codigo,nombre')}, follow=True)
        self.assertContains(r, 'debe ser Excel .xlsx')

    def test_existente_se_actualiza_sin_tocar_stock(self):
        with self.en_empresa():
            cat = Categoria.objects.create(nombre='R')
            p = Producto.objects.create(codigo='G1', nombre='Viejo', categoria=cat,
                                        unidad=Unidad.objects.get(abreviatura='UN'))
            servicios.ingresar(p, 10, Decimal('1'), self.dueno)
            r = importacion.validar(excel([{'codigo': 'g1', 'nombre': 'Nuevo nombre', 'categoria': 'R',
                                            'precio_venta_usd': 99, 'existencia_inicial': 50}]))
            self.assertTrue(r.valido)
            self.assertEqual(len(r.avisos), 1)                # se ignora la existencia
            importacion.aplicar(r, self.dueno)
            p = Producto.objects.con_stock().get(codigo='G1')
        self.assertEqual((p.nombre, p.precio_venta_usd, p.anot_stock_total), ('Nuevo nombre', Decimal('99'), 10))

    def test_importacion_queda_en_su_empresa(self):
        with self.en_empresa():
            importacion.aplicar(importacion.validar(excel([{'codigo': 'Z', 'nombre': 'Z', 'categoria': 'C'}])), self.dueno)
        with usando_empresa(self.otra):
            self.assertFalse(Producto.objects.exists())
            self.assertFalse(Categoria.objects.exists())

    def test_exportar_y_reimportar_solo_actualiza(self):
        with self.en_empresa():
            importacion.aplicar(importacion.validar(excel([
                {'codigo': 'A', 'nombre': 'A', 'categoria': 'C', 'marca': 'M', 'precio_venta_usd': 5}])), self.dueno)
        r = self.client.get(reverse('inventario:exportar'))
        with self.en_empresa():
            res = importacion.validar(r.content)
        self.assertTrue(res.valido, res.errores)
        self.assertEqual(res.resumen()['productos_actualizados'], 1)
        self.assertEqual(res.resumen()['productos_nuevos'], 0)


# ── Compras ───────────────────────────────────────────────────────────────────

class ComprasBase(Base):
    def setUp(self):
        super().setUp()
        with self.en_empresa():
            self.prov = Proveedor.objects.create(nombre='Distri')
            cat = Categoria.objects.create(nombre='R')
            un = Unidad.objects.get(abreviatura='UN')
            self.p1 = Producto.objects.create(codigo='P1', nombre='Uno', categoria=cat, unidad=un,
                                              proveedor_habitual=self.prov, precio_costo_usd=Decimal('10'),
                                              stock_minimo=5, factor_venta_dias=Decimal('1'))
            self.p2 = Producto.objects.create(codigo='P2', nombre='Dos', categoria=cat, unidad=un,
                                              maneja_lotes=True, precio_costo_usd=Decimal('3'),
                                              factor_venta_dias=Decimal('0'))


class ComprasTests(ComprasBase):
    def test_numeracion_por_empresa(self):
        with self.en_empresa():
            a = compras.crear_orden(self.prov, self.dueno)
            b = compras.crear_orden(self.prov, self.dueno)
        with usando_empresa(self.otra):
            c = compras.crear_orden(Proveedor.objects.create(nombre='X'), self.dueno)
        anio = HOY.year
        self.assertEqual([a.numero, b.numero, c.numero], [f'OC-{anio}-00001', f'OC-{anio}-00002', f'OC-{anio}-00001'])

    def test_flujo_completo_con_recepcion_parcial(self):
        r = self.client.post(reverse('compras:crear'), {'proveedor': self.prov.pk, 'notas': 'Urgente'})
        orden = OrdenCompra.todos.get()
        self.assertRedirects(r, reverse('compras:detalle', args=[orden.pk]))
        self.client.post(reverse('compras:agregar_item', args=[orden.pk]), {'producto': self.p1.pk, 'cantidad': 10, 'costo': '9'})
        self.client.post(reverse('compras:agregar_item', args=[orden.pk]), {'producto': self.p1.pk, 'cantidad': 2, 'costo': '9.5'})
        self.client.post(reverse('compras:agregar_item', args=[orden.pk]), {'producto': self.p2.pk, 'cantidad': 4, 'costo': '3'})
        items = {i.producto_id: i for i in ItemOrdenCompra.todos.filter(orden=orden)}
        self.assertEqual(items[self.p1.pk].cantidad_pedida, 12)       # mismo producto se suma
        self.client.post(reverse('compras:cambiar_estado', args=[orden.pk]), {'estado': 'ENVIADA'})
        self.assertEqual(OrdenCompra.todos.get().estado, 'ENVIADA')
        i1, i2 = items[self.p1.pk], items[self.p2.pk]
        # P2 exige lote: sin lote da error y no se recibe nada
        r = self.client.post(reverse('compras:recibir', args=[orden.pk]), {
            f'cantidad_{i1.pk}': 12, f'costo_{i1.pk}': '9.5', f'cantidad_{i2.pk}': 2, f'costo_{i2.pk}': '3'})
        self.assertContains(r, 'Indica el lote')
        self.assertFalse(Lote.todos.exists())
        self.client.post(reverse('compras:recibir', args=[orden.pk]), {
            f'cantidad_{i1.pk}': 12, f'costo_{i1.pk}': '9.5', f'cantidad_{i2.pk}': 2, f'lote_{i2.pk}': 'LX',
            f'costo_{i2.pk}': '3'})
        orden.refresh_from_db()
        self.assertEqual(orden.estado, 'PARCIAL')
        self.assertEqual(Lote.todos.get(producto=self.p2).cantidad_actual, 2)
        # Recibir de más es rechazado
        r = self.client.post(reverse('compras:recibir', args=[orden.pk]), {
            f'cantidad_{i2.pk}': 5, f'lote_{i2.pk}': 'LX', f'costo_{i2.pk}': '3'})
        self.assertContains(r, 'Entre 0 y 2')
        self.client.post(reverse('compras:recibir', args=[orden.pk]), {
            f'cantidad_{i2.pk}': 2, f'lote_{i2.pk}': 'LX', f'costo_{i2.pk}': '3'})
        orden.refresh_from_db()
        self.assertEqual(orden.estado, 'RECIBIDA')
        self.assertEqual(Lote.todos.get(producto=self.p2).cantidad_actual, 4)   # se sumó al mismo lote
        mov = MovimientoInventario.todos.filter(lote__producto=self.p1).get()
        self.assertIn(orden.numero, mov.motivo)

    def test_no_se_edita_orden_enviada(self):
        with self.en_empresa():
            orden = compras.crear_orden(self.prov, self.dueno, [(self.p1, 1, Decimal('1'))])
            compras.cambiar_estado(orden, 'ENVIADA')
            with self.assertRaises(compras.ErrorCompra):
                compras.agregar_item(orden, self.p2, 1, Decimal('1'))
            with self.assertRaises(compras.ErrorCompra):
                compras.cambiar_estado(orden, 'RECIBIDA')

    def test_no_se_envia_orden_vacia(self):
        with self.en_empresa():
            orden = compras.crear_orden(self.prov, self.dueno)
            with self.assertRaises(compras.ErrorCompra):
                compras.cambiar_estado(orden, 'ENVIADA')

    def test_pdf_orden(self):
        with self.en_empresa():
            orden = compras.crear_orden(self.prov, self.dueno, [(self.p1, 3, Decimal('10'))])
        for q in ('', '?costos=0'):
            r = self.client.get(reverse('compras:pdf', args=[orden.pk]) + q)
            self.assertTrue(r.content.startswith(b'%PDF'))

    def test_vendedor_no_accede_a_compras(self):
        self.client.force_login(self.vendedor)
        self.assertEqual(self.client.get(reverse('compras:lista')).status_code, 403)
        self.assertEqual(self.client.get(reverse('compras:panel_pedido')).status_code, 403)

    def test_orden_de_otra_empresa_da_404(self):
        with usando_empresa(self.otra):
            ajena = compras.crear_orden(Proveedor.objects.create(nombre='X'), self.dueno)
        self.assertEqual(self.client.get(reverse('compras:detalle', args=[ajena.pk])).status_code, 404)


class PanelPedidoTests(ComprasBase):
    def test_sugerido_por_estimado_y_por_ventas(self):
        with self.en_empresa():
            servicios.ingresar(self.p1, 4, Decimal('10'), self.dueno)
            grupos = sugerencias.calcular(dias_cobertura=30)
            s = grupos[self.prov][0]
            # factor 1 día/unidad → 30 para 30 días; hay 4 → pedir 26
            self.assertEqual((s.sugerido, s.fuente), (26, 'estimado'))
            # Con una orden abierta de 20, solo faltan 6
            compras.crear_orden(self.prov, self.dueno, [(self.p1, 20, Decimal('10'))])
            s = sugerencias.calcular(dias_cobertura=30)[self.prov][0]
            self.assertEqual(s.sugerido, 6)
            # P2 sin factor ni ventas y sin mínimo → no aparece
            self.assertFalse(any(x.producto == self.p2 for g in sugerencias.calcular(30).values() for x in g))

    def test_consumo_real_por_ventas(self):
        with self.en_empresa():
            lote = servicios.ingresar(self.p1, 100, Decimal('10'), self.dueno)
            servicios.apartar_fefo(self.p1, 90, self.dueno)
            servicios.descontar_apartado(lote, 90, self.dueno)       # 90 vendidas en 90 días → 1/día
            s = sugerencias.calcular(dias_cobertura=60)[self.prov][0]
        self.assertEqual((s.fuente, s.consumo_diario, s.disponible, s.sugerido), ('ventas', 1.0, 10, 50))

    def test_crear_orden_desde_panel(self):
        r = self.client.post(reverse('compras:crear_desde_panel'), {
            'proveedor': self.prov.pk, 'producto': [self.p1.pk], f'cantidad_{self.p1.pk}': 15})
        orden = OrdenCompra.todos.get()
        self.assertRedirects(r, reverse('compras:detalle', args=[orden.pk]))
        item = ItemOrdenCompra.todos.get(orden=orden)
        self.assertEqual((item.cantidad_pedida, item.costo_unitario_usd), (15, Decimal('10')))
        self.assertContains(self.client.get(reverse('compras:panel_pedido')), 'Panel de pedido')


class ListaPreciosTests(ComprasBase):
    def test_pdf_lista_de_precios(self):
        self.assertEqual(self.client.get(reverse('inventario:lista_precios')).status_code, 200)
        r = self.client.get(reverse('inventario:lista_precios'), {'generar': 1, 'mostrar_existencia': 1})
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.client.force_login(self.vendedor)     # el vendedor también puede generarla
        r = self.client.get(reverse('inventario:lista_precios'), {'generar': 1, 'con_existencia': 1})
        self.assertTrue(r.content.startswith(b'%PDF'))
