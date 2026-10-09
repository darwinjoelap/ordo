"""Pagos parciales (abonos), comisión al pago completo y estado de cuenta del cliente."""
from decimal import Decimal

from django.urls import reverse
from openpyxl import load_workbook
from io import BytesIO

from apps.comisiones import servicios as comisiones
from apps.comisiones.models import Comision
from apps.tasas import servicios as tasas
from apps.tasas.models import TasaCambio
from apps.ventas import servicios as ventas
from apps.ventas.models import Abono, Presupuesto

from .test_fase5 import Base


class AbonosConValidacionTests(Base):
    """La empresa exige validación: la venta va a «Por validar» solo cuando está pagada completa."""

    def setUp(self):
        super().setUp()
        self.perfil.requiere_validacion = True
        self.perfil.comision_requiere_pago = True
        self.perfil.save()

    def confirmada(self, n=5):                       # Glucosa $20 × 5 = 100 + IVA 16 = 116
        with self.empresa_ctx():
            p = ventas.crear(self.cliente, self.vendedor, self.empresa)
            ventas.agregar_item(p, self.prod, n, None, self.perfil, True)
            return ventas.confirmar(p, self.vendedor, self.empresa)

    def test_abonos_hasta_pagar(self):
        p = self.confirmada()
        self.assertEqual((p.estado, p.saldo_usd), ('POR_PAGAR', Decimal('116.00')))
        with self.empresa_ctx():
            p = ventas.registrar_abono(p, Decimal('50'), 'USD', 'ZELLE', self.vendedor)
            self.assertEqual((p.estado, p.pagado, p.abonado_usd, p.saldo_usd, p.abono_parcial),
                             ('POR_PAGAR', False, Decimal('50.00'), Decimal('66.00'), True))
            with self.assertRaisesRegex(ventas.ErrorVenta, 'supera'):
                ventas.registrar_abono(p, Decimal('70'), 'USD', 'ZELLE', self.vendedor)
            with self.assertRaisesRegex(ventas.ErrorVenta, 'validar|cobrar'):
                ventas.validar(p, self.dueno)
            p = ventas.registrar_abono(p, Decimal('66'), 'USD', 'PAGO_MOVIL', self.vendedor)
            self.assertEqual((p.estado, p.pagado, p.saldo_usd, p.metodo_pago), ('POR_VALIDAR', True, 0, 'MIXTO'))
            with self.assertRaisesRegex(ventas.ErrorVenta, 'pagada'):
                ventas.registrar_abono(p, Decimal('1'), 'USD', 'ZELLE', self.vendedor)
            p = ventas.anular_abono(Abono.objects.filter(presupuesto=p).last(), self.dueno)
            self.assertEqual((p.estado, p.pagado, p.saldo_usd), ('POR_PAGAR', False, Decimal('66.00')))

    def test_abono_en_bolivares_con_la_tasa(self):
        p = self.confirmada()
        with self.empresa_ctx():
            p = ventas.registrar_abono(p, Decimal('5000'), 'BS', 'PAGO_MOVIL', self.vendedor, tasa=Decimal('100'))
            a = Abono.objects.get(presupuesto=p)
            self.assertEqual((a.monto_usd, a.tasa_bs, p.saldo_usd), (Decimal('50.00'), Decimal('100'), Decimal('66.00')))
            tasas.registrar(Decimal('200'), TasaCambio.Fuente.MANUAL)
            p = ventas.registrar_abono(p, Decimal('13200'), 'BS', 'PAGO_MOVIL', self.vendedor)      # tasa del día
            self.assertTrue(p.pagado)

    def test_redondeo_de_un_centavo_cierra_la_venta(self):
        p = self.confirmada()
        with self.empresa_ctx():
            p = ventas.registrar_abono(p, Decimal('115.995'), 'USD', 'ZELLE', self.vendedor)
        self.assertTrue(p.pagado)

    def test_registrar_pago_completo_sigue_funcionando(self):
        p = self.confirmada()
        with self.empresa_ctx():
            p = ventas.registrar_abono(p, Decimal('16'), 'USD', 'ZELLE', self.vendedor)
            p = ventas.registrar_pago(p, True, 'ZELLE', usuario=self.dueno)
            self.assertEqual((p.pagado, p.abonado_usd, Abono.objects.filter(presupuesto=p).count()), (True, Decimal('116.00'), 2))
            p = ventas.registrar_pago(p, False)
            self.assertEqual((p.pagado, p.abonado_usd, p.estado), (False, 0, 'POR_PAGAR'))


class ComisionAlPagoCompletoTests(Base):
    """Sin validación: la venta se valida al confirmar; la comisión espera al pago COMPLETO."""

    def setUp(self):
        super().setUp()
        self.perfil.comision_requiere_pago = True
        self.perfil.save()

    def test_abono_parcial_no_libera_la_comision(self):
        p = self.venta([(self.prod, 5)])                  # validada, sin pagar
        with self.empresa_ctx():
            self.assertEqual(comisiones.saldo_disponible(self.perfil, self.vendedor), 0)
            p = ventas.registrar_abono(p, Decimal('100'), 'USD', 'ZELLE', self.vendedor)
            self.assertEqual(comisiones.saldo_disponible(self.perfil, self.vendedor), 0)
            ventas.registrar_abono(p, Decimal('16'), 'USD', 'ZELLE', self.vendedor)
            self.assertEqual(comisiones.saldo_disponible(self.perfil, self.vendedor), Comision.objects.get(presupuesto=p).monto_usd)

    def test_el_vendedor_ve_el_aviso(self):
        p = self.venta([(self.prod, 5)])
        with self.empresa_ctx():
            ventas.registrar_abono(p, Decimal('100'), 'USD', 'ZELLE', self.vendedor)
        self.client.force_login(self.vendedor)
        r = self.client.get(reverse('ventas:detalle', args=[p.pk]))
        self.assertContains(r, 'pagada por completo')
        self.assertContains(r, 'Los abonos parciales no la liberan')
        r = self.client.get(reverse('comisiones:detalle_vendedor'))
        self.assertContains(r, 'Espera pago completo')
        self.assertContains(r, '¿Cuándo se paga la comisión?')


class PantallaAbonosTests(Base):
    def setUp(self):
        super().setUp()
        self.p = self.venta([(self.prod, 5)])
        self.client.force_login(self.vendedor)

    def test_abonar_y_quitar(self):
        url = reverse('ventas:abonar', args=[self.p.pk])
        r = self.client.post(url, {'monto': '16', 'moneda': 'USD', 'metodo': 'ZELLE'}, follow=True)
        self.assertContains(r, 'Falta por cobrar')
        a = Abono.todos.get()
        self.client.post(reverse('ventas:anular_abono', args=[self.p.pk, a.pk]))
        self.assertTrue(Abono.todos.exists())                          # venta validada: el vendedor no puede quitarlo
        self.client.force_login(self.dueno)
        self.client.post(reverse('ventas:anular_abono', args=[self.p.pk, a.pk]))
        self.assertFalse(Abono.todos.exists())
        r = self.client.post(url, {'todo': '1', 'moneda': 'USD', 'metodo': 'ZELLE'}, follow=True)
        self.assertContains(r, 'pagada por completo')
        self.assertTrue(Presupuesto.todos.get(pk=self.p.pk).pagado)

    def test_datos_invalidos(self):
        url = reverse('ventas:abonar', args=[self.p.pk])
        self.assertContains(self.client.post(url, {'monto': 'abc', 'moneda': 'USD', 'metodo': 'ZELLE'}, follow=True), 'Revisa')
        self.assertContains(self.client.post(url, {'monto': '5', 'moneda': 'USD', 'metodo': ''}, follow=True), 'método')
        self.assertFalse(Abono.todos.exists())


class EstadoCuentaTests(Base):
    def setUp(self):
        super().setUp()
        self.pagada = self.venta([(self.prod, 1)])
        self.parcial = self.venta([(self.prod, 5)])
        self.debe = self.venta([(self.equipo, 1)])
        with self.empresa_ctx():
            ventas.registrar_pago(self.pagada, True, 'ZELLE', usuario=self.dueno)
            ventas.registrar_abono(self.parcial, Decimal('16'), 'USD', 'ZELLE', self.dueno)
            ventas.registrar_entrega(self.pagada, True)
            ventas.registrar_facturacion(self.pagada, True, 'F-001', 'C-1')
        self.client.force_login(self.dueno)
        self.url = reverse('clientes:detalle', args=[self.cliente.pk])

    def numeros(self, filtro):
        r = self.client.get(self.url, {'filtro': filtro})
        return {p.pk for p in r.context['pagina']}, r

    def test_filtros(self):
        self.assertEqual(self.numeros('pagadas')[0], {self.pagada.pk})
        self.assertEqual(self.numeros('abono_parcial')[0], {self.parcial.pk})
        self.assertEqual(self.numeros('por_cobrar')[0], {self.parcial.pk, self.debe.pk})
        self.assertEqual(self.numeros('facturadas')[0], {self.pagada.pk})
        self.assertEqual(self.numeros('entregadas')[0], {self.pagada.pk})
        self.assertEqual(self.numeros('por_entregar')[0], {self.parcial.pk, self.debe.pk})
        _, r = self.numeros('por_cobrar')
        self.assertEqual(r.context['totales']['saldo'], Decimal('100.00') + Decimal('232.00'))
        self.assertContains(r, 'Saldo por cobrar de este cliente')

    def test_reportes(self):
        r = self.client.get(reverse('clientes:estado_cuenta_pdf', args=[self.cliente.pk]), {'filtro': 'por_cobrar'})
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('Estado-cuenta', r['Content-Disposition'])
        r = self.client.get(reverse('clientes:estado_cuenta_excel', args=[self.cliente.pk]), {'filtro': 'abono_parcial'})
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws.cell(5, 1).value, self.parcial.numero)
        self.assertEqual(ws.cell(5, 8).value, 'Parcial')
        self.assertEqual(ws.cell(6, 1).value, 'TOTALES')

    def test_vendedor_solo_ve_sus_ventas(self):
        with self.empresa_ctx():
            otra = ventas.crear(self.cliente, self.vendedor2, self.empresa)
        self.client.force_login(self.vendedor)
        ids, _ = self.numeros('todos')
        self.assertNotIn(otra.pk, ids)
