"""
Exporta los productos y la existencia de BioLifeVentas al formato de la plantilla de Ordo.

SOLO LEE la base de datos de BioLifeVentas: no modifica nada.

Uso (PowerShell, desde la carpeta de BioLifeVentas y con SU venv):
    cd C:\\proyectos\\biolifeventas
    .\\venv\\Scripts\\Activate.ps1
    pip install openpyxl                      # solo en el venv, no se agrega a requirements
    # Para leer PRODUCCIÓN, apunta a la BD de Railway (URL pública de Postgres):
    $env:DATABASE_URL = "postgresql://..."
    python manage.py shell -c "exec(open(r'C:\\proyectos\\ordo\\scripts\\biolifeventas_a_ordo.py', encoding='utf-8').read())"

Genera: productos_biolife_para_ordo.xlsx en la carpeta actual.

«dias_por_unidad» se CALCULA con las ventas reales de BioLifeVentas (no se copia factor_venta_dias):
el cron `recalcular_factores_venta` de BioLifeVentas guarda ahí los días que dura el stock, no los días
por unidad vendida, y copiarlo haría creer a Ordo que los productos casi no se venden.
    con ventas en 90 días  → días de historial / unidades vendidas (historial = desde la 1.ª venta, 30 a 90 días)
    sin ventas en 90 días  → 90 (rota menos de 1 unidad cada 90 días), salvo factor manual (auto-ajuste apagado)
"""
import os
from collections import Counter
from datetime import timedelta

from django.db.models import Min, Sum
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from inventario.models import Lote, MovimientoInventario, Producto  # modelos de BioLifeVentas

# Mismo orden de columnas que la plantilla de Ordo (apps/inventario/importacion.py)
COLUMNAS = ['codigo', 'nombre', 'categoria', 'subcategoria', 'marca', 'unidad', 'proveedor', 'costo_usd',
            'precio_venta_usd', 'stock_minimo', 'dias_por_unidad', 'maneja_lote', 'maneja_vencimiento',
            'descripcion', 'activo', 'existencia_inicial', 'numero_lote', 'fecha_vencimiento']

# Unidades de BioLifeVentas → abreviatura en Ordo (todas existen en Ordo desde inventario 0006; la mayoría es VIAL)
UNIDADES = {'UN': 'UN', 'CAJA': 'CAJA', 'KIT': 'KIT', 'FRASCO': 'FCO',
            'VIAL': 'VIAL', 'ML': 'ML', 'TEST': 'TEST', 'TUBO': 'TUBO', 'ROLLO': 'ROLLO'}
UNIDADES_ORDO_INICIALES = {'UN', 'CAJA', 'KIT', 'PAQ', 'FCO', 'KG', 'L', 'M', 'VIAL', 'ML', 'TEST', 'TUBO', 'ROLLO'}

si_no = lambda b: 'SI' if b else 'NO'  # noqa: E731
salida = os.environ.get('ORDO_SALIDA', 'productos_biolife_para_ordo.xlsx')

wb = Workbook()
ws = wb.active
ws.title = 'Productos'
ws.append(COLUMNAS)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='053D74')

productos = (Producto.objects.select_related('categoria', 'subcategoria', 'marca', 'proveedor_habitual')
             .order_by('codigo'))
lotes_por_producto = {}
for l in Lote.objects.filter(cantidad_actual__gt=0).order_by('fecha_vencimiento', 'fecha_ingreso'):
    lotes_por_producto.setdefault(l.producto_id, []).append(l)

# ── Consumo real de los últimos 90 días ──────────────────────────────────────
DIAS, MINIMO = 90, 30
ahora = timezone.now()
vendidas_90 = {r['lote__producto']: abs(r['t'] or 0) for r in
               MovimientoInventario.objects.filter(tipo='VENTA', fecha__gte=ahora - timedelta(days=DIAS))
               .values('lote__producto').annotate(t=Sum('cantidad'))}
primera_venta = dict(MovimientoInventario.objects.filter(tipo='VENTA').values('lote__producto')
                     .annotate(f=Min('fecha')).values_list('lote__producto', 'f'))
origen_factor = Counter()


def dias_por_unidad(p):
    vendidas = vendidas_90.get(p.pk, 0)
    if vendidas > 0:
        antiguedad = (ahora - primera_venta[p.pk]).days if p.pk in primera_venta else DIAS
        base = min(DIAS, max(MINIMO, antiguedad))
        origen_factor['ventas'] += 1
        return max(0.01, round(base / vendidas, 2))
    if not getattr(p, 'factor_auto_ajuste', True) and p.factor_venta_dias and p.factor_venta_dias > 0:
        origen_factor['manual'] += 1
        return float(p.factor_venta_dias)
    origen_factor['sin_ventas'] += 1
    return float(DIAS)


unidades_usadas = Counter()
n_prod = n_lotes = unidades_totales = apartadas = 0
sin_lote_con_exigencia = []
for p in productos:
    unidad = UNIDADES.get(p.unidad, p.unidad)
    unidades_usadas[unidad] += 1
    base = [p.codigo, p.nombre, p.categoria.nombre, p.subcategoria.nombre if p.subcategoria else None,
            p.marca.nombre if p.marca else None, unidad,
            p.proveedor_habitual.nombre if p.proveedor_habitual else None,
            float(p.precio_costo_usd), float(p.precio_venta_usd), p.stock_minimo,
            dias_por_unidad(p), si_no(p.requiere_lote), si_no(p.requiere_vencimiento),
            p.descripcion or None, si_no(p.activo)]
    lotes = lotes_por_producto.get(p.pk, [])
    n_prod += 1
    if not lotes:
        ws.append(base + [None, None, None])
        continue
    for i, l in enumerate(lotes):
        numero = l.numero_lote or (f'SIN-LOTE-{i + 1}' if p.requiere_lote else None)
        if p.requiere_lote and not l.numero_lote:
            sin_lote_con_exigencia.append(p.codigo)
        fila = base if i == 0 else [p.codigo] + [None] * 14
        ws.append(fila + [l.cantidad_actual, numero, l.fecha_vencimiento])
        ws.cell(row=ws.max_row, column=18).number_format = 'DD/MM/YYYY'
        n_lotes += 1
        unidades_totales += l.cantidad_actual
        apartadas += l.cantidad_apartada

for i, ancho in enumerate([14, 34, 16, 18, 14, 8, 20, 10, 12, 10, 12, 10, 12, 20, 8, 12, 14, 14], start=1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

faltan = sorted(u for u in unidades_usadas if u not in UNIDADES_ORDO_INICIALES)
wn = wb.create_sheet('Notas migración')
wn.append(['Resumen de la exportación desde BioLifeVentas'])
wn['A1'].font = Font(bold=True, size=13)
wn.append([f'Productos: {n_prod} · Lotes con existencia: {n_lotes} · Unidades en existencia: {unidades_totales}'])
wn.append([f'Unidades APARTADAS hoy en BioLifeVentas: {apartadas} (entran como existencia normal; '
           'los apartados se migran con los presupuestos en la fase 7).'])
wn.append([f'Días por unidad: {origen_factor["ventas"]} productos calculados con ventas reales de 90 días, '
           f'{origen_factor["manual"]} con factor manual, {origen_factor["sin_ventas"]} sin ventas recientes (=90).'])
wn.append([])
if faltan:
    wn.append(['ANTES DE IMPORTAR crea en Ordo › Inventario › Catálogo › Unidades estas abreviaturas:'])
    for u in faltan:
        wn.append([f'   {u}  ({unidades_usadas[u]} productos)'])
else:
    wn.append(['Todas las unidades ya existen en Ordo.'])
if sin_lote_con_exigencia:
    wn.append([])
    wn.append([f'{len(set(sin_lote_con_exigencia))} productos exigen lote pero tenían lotes sin número: '
               'se les puso SIN-LOTE-n. Revísalos.'])
wn.column_dimensions['A'].width = 110
wb.save(salida)

print(f'Listo: {os.path.abspath(salida)}')
print(f'  {n_prod} productos · {n_lotes} lotes · {unidades_totales} unidades ({apartadas} apartadas)')
print(f'  Días por unidad: {origen_factor["ventas"]} con ventas reales · {origen_factor["manual"]} manual · '
      f'{origen_factor["sin_ventas"]} sin ventas recientes')
if faltan:
    print(f'  Crea primero en Ordo estas unidades: {", ".join(faltan)}')
