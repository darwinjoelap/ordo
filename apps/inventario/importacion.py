"""
Importación masiva de productos desde Excel (.xlsx).

Flujo:
  1. El usuario descarga la plantilla (generar_plantilla) — trae instrucciones y ejemplos.
  2. Sube el archivo → validar() revisa TODO sin guardar nada y arma un resumen.
  3. Si no hay errores, confirma → aplicar() crea/actualiza en una sola transacción.

Las columnas están definidas UNA sola vez en COLUMNAS: de ahí salen la plantilla,
la hoja de instrucciones, la validación y la guía de la documentación.
"""
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.proveedores.models import Proveedor

from . import servicios
from .models import Categoria, Marca, Producto, Subcategoria, Unidad

MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024
HOJA = 'Productos'


@dataclass(frozen=True)
class Columna:
    clave: str
    titulo: str
    obligatoria: bool
    formato: str
    ejemplo: object
    nota: str
    ancho: int = 16


COLUMNAS = [
    Columna('codigo', 'codigo', True, 'Texto, máx. 50', 'GLU-001',
            'Único en tu empresa. Se guarda en MAYÚSCULAS. Si ya existe, el producto se ACTUALIZA.', 14),
    Columna('nombre', 'nombre', True, 'Texto, máx. 200', 'Glucosa enzimática 500 ml', 'Nombre que verá el cliente.', 34),
    Columna('categoria', 'categoria', True, 'Texto', 'Reactivos', 'Si no existe, se crea.', 16),
    Columna('subcategoria', 'subcategoria', False, 'Texto', 'Química sanguínea',
            'Opcional. Se crea dentro de la categoría si no existe.', 18),
    Columna('marca', 'marca', False, 'Texto', 'Wiener', 'Opcional. Si no existe, se crea.', 14),
    Columna('unidad', 'unidad', False, 'Abreviatura', 'KIT',
            'Abreviatura de una unidad existente (ver hoja Unidades). Vacío = UN.', 10),
    Columna('proveedor', 'proveedor', False, 'Texto', 'Distribuidora Lara',
            'Proveedor habitual. Si no existe, se crea.', 20),
    Columna('costo_usd', 'costo_usd', False, 'Número ≥ 0 (punto o coma decimal)', 42.5, 'Costo unitario en USD. Vacío = 0.', 12),
    Columna('precio_venta_usd', 'precio_venta_usd', False, 'Número ≥ 0', 86, 'Precio de venta en USD. Vacío = 0.', 16),
    Columna('stock_minimo', 'stock_minimo', False, 'Entero ≥ 0', 5, 'Alerta de "bajo mínimo". Vacío = 0.', 13),
    Columna('dias_por_unidad', 'dias_por_unidad', False, 'Número > 0', 7,
            'Rotación estimada: cada cuántos días se vende 1 unidad (usado por el panel de pedido). Vacío = 30.', 15),
    Columna('maneja_lote', 'maneja_lote', False, 'SI / NO', 'SI', 'Exige número de lote en cada ingreso. Vacío = NO.', 13),
    Columna('maneja_vencimiento', 'maneja_vencimiento', False, 'SI / NO', 'SI',
            'Exige fecha de vencimiento en cada ingreso. Vacío = NO.', 18),
    Columna('descripcion', 'descripcion', False, 'Texto', '', 'Opcional.', 24),
    Columna('activo', 'activo', False, 'SI / NO', 'SI', 'Vacío = SI.', 8),
    Columna('existencia_inicial', 'existencia_inicial', False, 'Entero ≥ 0', 12,
            'Solo para productos NUEVOS. Crea un ingreso inicial con el costo_usd.', 16),
    Columna('numero_lote', 'numero_lote', False, 'Texto', 'L2601',
            'Lote de la existencia inicial. Obligatorio si maneja_lote = SI y hay existencia.', 13),
    Columna('fecha_vencimiento', 'fecha_vencimiento', False, 'Fecha dd/mm/aaaa', date(2027, 3, 31),
            'Vencimiento del lote inicial. Obligatoria si maneja_vencimiento = SI y hay existencia.', 16),
]
CLAVES = [c.clave for c in COLUMNAS]
OBLIGATORIAS = [c.clave for c in COLUMNAS if c.obligatoria]

REGLAS = [
    'La primera fila de la hoja "Productos" son los títulos: no los cambies ni cambies su orden.',
    'Una fila por producto. Para cargar VARIOS LOTES del mismo producto, repite el código en otra fila '
    'llenando solo codigo, existencia_inicial, numero_lote y fecha_vencimiento.',
    'Si el código ya existe en Ordo, se actualizan sus datos (nombre, precios, categoría…). '
    'La existencia NO se toca: para eso usa Ingreso o Ajuste.',
    'Las filas totalmente vacías se ignoran. Máximo 5.000 filas y 5 MB por archivo.',
    'Antes de guardar nada, Ordo valida todo el archivo y te muestra los errores por fila. '
    'Si hay un solo error, no se importa nada: corrige y vuelve a subirlo.',
    'Los montos van en dólares (USD). Se aceptan 12.50 o 12,50; no uses símbolos de moneda.',
]


# ── Plantilla ─────────────────────────────────────────────────────────────────

AZUL = '053D74'
VERDE = '11988D'


def generar_plantilla(unidades=None, con_ejemplos=True):
    """Excel con hojas Productos, Instrucciones y Unidades."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    titulo_obl = PatternFill('solid', fgColor=AZUL)
    titulo_opc = PatternFill('solid', fgColor=VERDE)
    for i, col in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i, value=col.titulo)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = titulo_obl if col.obligatoria else titulo_opc
        celda.alignment = Alignment(horizontal='center')
        celda.comment = Comment(f'{"OBLIGATORIA" if col.obligatoria else "Opcional"} · {col.formato}\n{col.nota}', 'Ordo')
        ws.column_dimensions[get_column_letter(i)].width = col.ancho
    ws.freeze_panes = 'B2'

    if con_ejemplos:
        ejemplos = [
            {c.clave: c.ejemplo for c in COLUMNAS},
            {'codigo': 'GLU-001', 'existencia_inicial': 8, 'numero_lote': 'L2602', 'fecha_vencimiento': date(2027, 6, 30)},
            {'codigo': 'TUB-EDTA', 'nombre': 'Tubo EDTA 4 ml (caja x100)', 'categoria': 'Consumibles', 'marca': 'BD',
             'unidad': 'CAJA', 'costo_usd': 18, 'precio_venta_usd': 32, 'stock_minimo': 10,
             'maneja_lote': 'NO', 'maneja_vencimiento': 'NO', 'existencia_inicial': 40},
        ]
        for r, ej in enumerate(ejemplos, start=2):
            for i, clave in enumerate(CLAVES, start=1):
                valor = ej.get(clave, '')
                celda = ws.cell(row=r, column=i, value=valor if valor != '' else None)
                if isinstance(valor, date):
                    celda.number_format = 'DD/MM/YYYY'

    for i, col in enumerate(COLUMNAS, start=1):
        if col.formato == 'SI / NO':
            dv = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}{MAX_FILAS + 1}')
            ws.add_data_validation(dv)
        if col.clave == 'fecha_vencimiento':
            for fila in range(2, 200):
                ws.cell(row=fila, column=i).number_format = 'DD/MM/YYYY'

    # Instrucciones
    wi = wb.create_sheet('Instrucciones')
    wi['A1'] = 'Cómo llenar la plantilla de productos de Ordo'
    wi['A1'].font = Font(bold=True, size=14, color=AZUL)
    fila = 3
    for regla in REGLAS:
        wi.cell(row=fila, column=1, value=f'• {regla}').alignment = Alignment(wrap_text=True, vertical='top')
        wi.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
        wi.row_dimensions[fila].height = 32
        fila += 1
    fila += 1
    for i, t in enumerate(['Columna', '¿Obligatoria?', 'Formato', 'Ejemplo', 'Notas'], start=1):
        c = wi.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
    for col in COLUMNAS:
        fila += 1
        ejemplo = col.ejemplo.strftime('%d/%m/%Y') if isinstance(col.ejemplo, date) else col.ejemplo
        for i, v in enumerate([col.titulo, 'SÍ' if col.obligatoria else 'no', col.formato, ejemplo, col.nota], start=1):
            wi.cell(row=fila, column=i, value=v).alignment = Alignment(wrap_text=True, vertical='top')
    for letra, ancho in zip('ABCDE', (20, 13, 26, 26, 70)):
        wi.column_dimensions[letra].width = ancho

    # Unidades válidas
    wu = wb.create_sheet('Unidades')
    wu.append(['abreviatura', 'nombre'])
    for c in wu[1]:
        c.font = Font(bold=True)
    for abrev, nombre in (unidades or [('UN', 'Unidad'), ('CAJA', 'Caja'), ('KIT', 'Kit')]):
        wu.append([abrev, nombre])
    wu.column_dimensions['A'].width = 14
    wu.column_dimensions['B'].width = 24

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ── Lectura y validación ──────────────────────────────────────────────────────

class ErrorArchivo(Exception):
    """El archivo no se puede leer (formato, hoja, títulos)."""


@dataclass
class Fila:
    numero: int
    datos: dict


@dataclass
class ProductoImportado:
    codigo: str
    fila: int
    datos: dict
    lotes: list = field(default_factory=list)   # [(fila, cantidad, numero_lote, fecha)]
    existente: Producto | None = None


@dataclass
class Resultado:
    productos: 'OrderedDict[str, ProductoImportado]'
    errores: list          # [(fila, columna, mensaje)]
    avisos: list           # [(fila, mensaje)]
    nuevas: dict           # {'categorias': set, 'subcategorias': set, 'marcas': set, 'proveedores': set}

    @property
    def valido(self):
        return not self.errores

    def resumen(self):
        nuevos = [p for p in self.productos.values() if not p.existente]
        return {
            'productos_nuevos': len(nuevos),
            'productos_actualizados': len(self.productos) - len(nuevos),
            'lotes': sum(len(p.lotes) for p in nuevos),
            'unidades_iniciales': sum(c for p in nuevos for _, c, _, _ in p.lotes),
            'categorias_nuevas': sorted(self.nuevas['categorias']),
            'subcategorias_nuevas': sorted(f'{c} › {s}' for c, s in self.nuevas['subcategorias']),
            'marcas_nuevas': sorted(self.nuevas['marcas']),
            'proveedores_nuevos': sorted(self.nuevas['proveedores']),
            'errores': len(self.errores),
            'avisos': len(self.avisos),
        }


def _texto(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _si_no(v, defecto):
    t = _texto(v).upper().replace('Í', 'I')
    if t == '':
        return defecto
    if t in ('SI', 'S', 'TRUE', '1', 'X', 'VERDADERO'):
        return True
    if t in ('NO', 'N', 'FALSE', '0', 'FALSO'):
        return False
    raise ValueError('usa SI o NO')


def _decimal(v, defecto=Decimal('0')):
    t = _texto(v).replace('$', '').replace(' ', '')
    if t == '':
        return defecto
    if ',' in t and '.' in t:          # 1.234,56 → 1234.56
        t = t.replace('.', '').replace(',', '.')
    else:
        t = t.replace(',', '.')
    try:
        d = Decimal(t).quantize(Decimal('0.01'))
    except InvalidOperation:
        raise ValueError('debe ser un número')
    if d < 0:
        raise ValueError('no puede ser negativo')
    return d


def _entero(v, defecto=0):
    t = _texto(v)
    if t == '':
        return defecto
    try:
        d = Decimal(t.replace(',', '.'))
    except InvalidOperation:
        raise ValueError('debe ser un número entero')
    if d != d.to_integral_value() or d < 0:
        raise ValueError('debe ser un entero ≥ 0')
    return int(d)


def _fecha(v):
    if v in (None, ''):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _texto(v)
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise ValueError('usa el formato dd/mm/aaaa')


def leer_filas(contenido):
    if len(contenido) > MAX_BYTES:
        raise ErrorArchivo('El archivo supera 5 MB.')
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorArchivo('No se pudo leer el archivo. Debe ser un Excel .xlsx (usa la plantilla de Ordo).')
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    try:
        encabezado = [(_texto(c).lower()) for c in next(filas)]
    except StopIteration:
        raise ErrorArchivo('La hoja está vacía.')
    faltan = [c for c in OBLIGATORIAS if c not in encabezado]
    if faltan:
        raise ErrorArchivo(f'Faltan columnas obligatorias: {", ".join(faltan)}. Usa la plantilla de Ordo.')
    indices = {clave: encabezado.index(clave) for clave in CLAVES if clave in encabezado}
    resultado = []
    for numero, valores in enumerate(filas, start=2):
        if valores is None or all(v in (None, '') for v in valores):
            continue
        datos = {clave: (valores[i] if i < len(valores) else None) for clave, i in indices.items()}
        resultado.append(Fila(numero, datos))
        if len(resultado) > MAX_FILAS:
            raise ErrorArchivo(f'El archivo tiene más de {MAX_FILAS} filas con datos.')
    wb.close()
    if not resultado:
        raise ErrorArchivo('El archivo no tiene filas con datos.')
    return resultado


def validar(contenido):
    """Valida el archivo completo sin guardar nada."""
    filas = leer_filas(contenido)
    errores, avisos = [], []
    productos = OrderedDict()
    nuevas = {'categorias': set(), 'subcategorias': set(), 'marcas': set(), 'proveedores': set()}

    unidades = {u.abreviatura.upper(): u for u in Unidad.objects.all()}
    categorias = {c.nombre.lower(): c for c in Categoria.objects.all()}
    subcategorias = {(s.categoria.nombre.lower(), s.nombre.lower()) for s in Subcategoria.objects.select_related('categoria')}
    marcas = {m.nombre.lower() for m in Marca.objects.all()}
    proveedores = {p.nombre.lower() for p in Proveedor.objects.all()}
    existentes = {p.codigo: p for p in Producto.objects.all()}

    def error(fila, columna, mensaje):
        errores.append((fila, columna, mensaje))

    for f in filas:
        d = f.datos
        codigo = _texto(d.get('codigo')).upper()
        if not codigo:
            error(f.numero, 'codigo', 'es obligatorio')
            continue
        if len(codigo) > 50:
            error(f.numero, 'codigo', 'máximo 50 caracteres')
            continue

        existencia_txt = _texto(d.get('existencia_inicial'))
        if codigo in productos:
            # Fila adicional: solo agrega un lote al producto ya definido
            p = productos[codigo]
            otros = [c for c in CLAVES if c not in ('codigo', 'existencia_inicial', 'numero_lote', 'fecha_vencimiento')
                     and _texto(d.get(c))]
            if otros:
                avisos.append((f.numero, f'Código repetido: se ignoran las columnas {", ".join(otros)} '
                                         f'(se usan las de la fila {p.fila}).'))
            _leer_lote(f, d, p, error)
            continue

        nombre = _texto(d.get('nombre'))
        categoria = _texto(d.get('categoria'))
        if not nombre:
            error(f.numero, 'nombre', 'es obligatorio')
        elif len(nombre) > 200:
            error(f.numero, 'nombre', 'máximo 200 caracteres')
        if not categoria:
            error(f.numero, 'categoria', 'es obligatoria')

        datos = {'nombre': nombre, 'categoria': categoria, 'subcategoria': _texto(d.get('subcategoria')),
                 'marca': _texto(d.get('marca')), 'proveedor': _texto(d.get('proveedor')),
                 'descripcion': _texto(d.get('descripcion'))}
        abrev = _texto(d.get('unidad')).upper() or 'UN'
        if abrev not in unidades:
            error(f.numero, 'unidad', f'"{abrev}" no existe. Válidas: {", ".join(sorted(unidades))}')
        else:
            datos['unidad'] = unidades[abrev]
        for clave, conversor, defecto in (
                ('costo_usd', _decimal, Decimal('0')), ('precio_venta_usd', _decimal, Decimal('0')),
                ('stock_minimo', _entero, 0), ('dias_por_unidad', _decimal, Decimal('30')),
                ('maneja_lote', _si_no, False),
                ('maneja_vencimiento', _si_no, False), ('activo', _si_no, True)):
            try:
                datos[clave] = conversor(d.get(clave), defecto)
            except ValueError as e:
                error(f.numero, clave, str(e))

        if datos.get('dias_por_unidad') == 0:
            error(f.numero, 'dias_por_unidad', 'debe ser mayor que 0')
        p = ProductoImportado(codigo=codigo, fila=f.numero, datos=datos, existente=existentes.get(codigo))
        productos[codigo] = p
        _leer_lote(f, d, p, error)

        if categoria and categoria.lower() not in categorias:
            nuevas['categorias'].add(categoria)
        if datos['subcategoria'] and (categoria.lower(), datos['subcategoria'].lower()) not in subcategorias:
            nuevas['subcategorias'].add((categoria, datos['subcategoria']))
        if datos['marca'] and datos['marca'].lower() not in marcas:
            nuevas['marcas'].add(datos['marca'])
        if datos['proveedor'] and datos['proveedor'].lower() not in proveedores:
            nuevas['proveedores'].add(datos['proveedor'])

    # Reglas que dependen del producto completo
    for p in productos.values():
        if p.existente and p.lotes:
            avisos.append((p.fila, f'{p.codigo} ya existe: se actualizan sus datos pero se IGNORA la existencia '
                                   f'inicial ({sum(c for _, c, _, _ in p.lotes)} u.). Usa Ingreso para sumar stock.'))
            p.lotes = []
        lotes_vistos = set()
        for fila, cantidad, lote, venc in p.lotes:
            if p.datos.get('maneja_lote') and not lote:
                error(fila, 'numero_lote', f'{p.codigo} maneja lote: indica el número de lote')
            if p.datos.get('maneja_vencimiento') and not venc:
                error(fila, 'fecha_vencimiento', f'{p.codigo} maneja vencimiento: indica la fecha')
            if lote and lote in lotes_vistos:
                error(fila, 'numero_lote', f'el lote {lote} está repetido para {p.codigo}')
            lotes_vistos.add(lote)

    return Resultado(productos, sorted(errores), sorted(avisos), nuevas)


def _leer_lote(f, d, p, error):
    try:
        cantidad = _entero(d.get('existencia_inicial'), 0)
    except ValueError as e:
        error(f.numero, 'existencia_inicial', str(e))
        return
    lote = _texto(d.get('numero_lote'))[:50]
    try:
        venc = _fecha(d.get('fecha_vencimiento'))
    except ValueError as e:
        error(f.numero, 'fecha_vencimiento', str(e))
        return
    if cantidad > 0:
        p.lotes.append((f.numero, cantidad, lote, venc))
    elif lote or venc:
        error(f.numero, 'existencia_inicial', 'indica la cantidad del lote (o borra lote y vencimiento)')


# ── Aplicación ────────────────────────────────────────────────────────────────

@transaction.atomic
def aplicar(resultado, usuario):
    """Crea/actualiza todo. Si algo falla, no queda nada a medias."""
    if not resultado.valido:
        raise ValueError('El archivo tiene errores.')
    cache_cat = {c.nombre.lower(): c for c in Categoria.objects.all()}
    cache_sub = {(s.categoria_id, s.nombre.lower()): s for s in Subcategoria.objects.all()}
    cache_marca = {m.nombre.lower(): m for m in Marca.objects.all()}
    cache_prov = {p.nombre.lower(): p for p in Proveedor.objects.all()}

    def obtener(cache, clave, crear):
        if clave not in cache:
            cache[clave] = crear()
        return cache[clave]

    creados = actualizados = lotes = 0
    for p in resultado.productos.values():
        d = p.datos
        cat = obtener(cache_cat, d['categoria'].lower(), lambda: Categoria.objects.create(nombre=d['categoria']))
        sub = None
        if d['subcategoria']:
            sub = obtener(cache_sub, (cat.pk, d['subcategoria'].lower()),
                          lambda: Subcategoria.objects.create(categoria=cat, nombre=d['subcategoria']))
        marca = obtener(cache_marca, d['marca'].lower(), lambda: Marca.objects.create(nombre=d['marca'])) if d['marca'] else None
        prov = obtener(cache_prov, d['proveedor'].lower(),
                       lambda: Proveedor.objects.create(nombre=d['proveedor'])) if d['proveedor'] else None
        campos = dict(nombre=d['nombre'], descripcion=d['descripcion'], categoria=cat, subcategoria=sub, marca=marca,
                      unidad=d['unidad'], proveedor_habitual=prov, precio_costo_usd=d['costo_usd'],
                      precio_venta_usd=d['precio_venta_usd'], stock_minimo=d['stock_minimo'],
                      factor_venta_dias=d['dias_por_unidad'],
                      maneja_lotes=d['maneja_lote'], maneja_vencimiento=d['maneja_vencimiento'], activo=d['activo'])
        if p.existente:
            for k, v in campos.items():
                setattr(p.existente, k, v)
            p.existente.save()
            actualizados += 1
            continue
        producto = Producto.objects.create(codigo=p.codigo, **campos)
        creados += 1
        for _, cantidad, lote, venc in p.lotes:
            servicios.ingresar(producto, cantidad, d['costo_usd'], usuario, numero_lote=lote,
                               fecha_vencimiento=venc, proveedor=prov, notas='Existencia inicial (importación)')
            lotes += 1
    return {'creados': creados, 'actualizados': actualizados, 'lotes': lotes}


def exportar_productos(productos):
    """Exporta productos con el MISMO formato de la plantilla (sirve para editar en Excel y reimportar)."""
    contenido = generar_plantilla(
        unidades=[(u.abreviatura, u.nombre) for u in Unidad.objects.all()], con_ejemplos=False)
    wb = load_workbook(BytesIO(contenido))
    ws = wb[HOJA]
    si_no = lambda b: 'SI' if b else 'NO'  # noqa: E731
    for p in productos:
        ws.append([p.codigo, p.nombre, p.categoria.nombre, p.subcategoria.nombre if p.subcategoria else None,
                   p.marca.nombre if p.marca else None, p.unidad.abreviatura,
                   p.proveedor_habitual.nombre if p.proveedor_habitual else None,
                   float(p.precio_costo_usd), float(p.precio_venta_usd), p.stock_minimo, float(p.factor_venta_dias),
                   si_no(p.maneja_lotes),
                   si_no(p.maneja_vencimiento), p.descripcion or None, si_no(p.activo)])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
