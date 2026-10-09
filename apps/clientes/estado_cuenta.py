"""
Estado de cuenta de un cliente: sus documentos filtrados por situación (por cobrar, abonadas, pagadas,
por entregar, facturadas…), con totales, en pantalla, PDF y Excel.
"""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.db.models import F, Q
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import Spacer, Table, TableStyle

from apps.core.pdf import ESTILO_CELDA, Paragraph, documento, encabezado_empresa, pie_empresa
from apps.core.templatetags.ordo import usd
from apps.ventas.models import Presupuesto

E = Presupuesto.Estado
CONF = Presupuesto.CONFIRMADAS
CERO = Decimal('0')

# clave: (etiqueta, condición). El orden es el de los botones en pantalla.
FILTROS = {
    'todos': ('Todos', ~Q(estado=E.BORRADOR)),
    'por_cobrar': ('Por cobrar', Q(estado=E.POR_PAGAR) | Q(estado=E.VALIDADA, pagado=False)),
    'abono_parcial': ('Abono parcial', Q(estado__in=CONF, pagado=False, abonado_usd__gt=0)),
    'pagadas': ('Pagadas', Q(estado__in=(*CONF, E.DEVUELTA), pagado=True)),
    'por_entregar': ('Por entregar', Q(estado__in=CONF, entregado=False)),
    'entregadas': ('Entregadas', Q(estado__in=(*CONF, E.DEVUELTA), entregado=True)),
    'facturadas': ('Facturadas', Q(facturado=True)),
    'sin_facturar': ('Sin facturar', Q(estado=E.VALIDADA, facturado=False)),
    'presupuestos': ('Presupuestos abiertos', Q(estado__in=(E.EMITIDO, E.APARTADO))),
    'devoluciones': ('Con devolución', Q(devuelto_usd__gt=0)),
    'anulados': ('Cancelados / vencidos', Q(estado__in=(E.CANCELADO, E.VENCIDO, E.RECHAZADA))),
}


def documentos(qs, filtro, desde=None, hasta=None):
    """qs: presupuestos ya limitados al cliente (y a lo que el usuario puede ver)."""
    filtro = filtro if filtro in FILTROS else 'todos'
    qs = qs.filter(FILTROS[filtro][1])
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    return qs.select_related('vendedor').order_by('-fecha', '-pk'), filtro


def totales(docs):
    t = {'n': 0, 'total': CERO, 'abonado': CERO, 'saldo': CERO}
    for p in docs:
        t['n'] += 1
        t['total'] += p.neto_usd
        if p.confirmada or p.es_venta:
            t['abonado'] += p.abonado_usd
            t['saldo'] += p.saldo_usd
    return t


def conteos(qs):
    """Cuántos documentos hay en cada filtro (para los botones)."""
    return {k: qs.filter(cond).count() for k, (_, cond) in FILTROS.items()}


def _situacion(p):
    """Texto corto de cobro, entrega y factura de un documento."""
    partes = [p.get_estado_display()]
    if p.confirmada or p.es_venta:
        partes.append('Pagada' if p.pagado else ('Abono parcial' if p.abonado_usd else 'Sin pagos'))
        partes.append(f'Entregada {p.fecha_entrega:%d/%m/%Y}' if p.entregado and p.fecha_entrega else
                      ('Entregada' if p.entregado else 'Por entregar'))
    if p.facturado:
        partes.append(f'Fact. {p.numero_factura}')
    return partes


def _rango(desde, hasta):
    if desde and hasta:
        return f'del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'
    if desde:
        return f'desde el {desde:%d/%m/%Y}'
    if hasta:
        return f'hasta el {hasta:%d/%m/%Y}'
    return 'todas las fechas'


def pdf(cliente, docs, filtro, empresa, desde=None, hasta=None):
    perfil = empresa.perfil
    color = colors.HexColor(perfil.color_principal or '#053D74')
    buffer = BytesIO()
    doc = documento(buffer, titulo_pdf=f'Estado de cuenta {cliente.nombre}', horizontal=True)
    hoy = datetime.now().strftime('%d/%m/%Y')
    e = encabezado_empresa(empresa, titulo='ESTADO DE CUENTA',
                           subtitulo=f'{FILTROS[filtro][0]} · {_rango(desde, hasta)}<br/>Emitido: {hoy}')
    e += [Paragraph(f'<b>{cliente.nombre}</b> · RIF/Cédula: {cliente.rif or "—"}'
                    + (f' · Tel: {cliente.telefono}' if cliente.telefono else ''), ESTILO_CELDA), Spacer(1, 3 * mm)]
    cab = ['N°', 'Fecha', 'Vendedor', 'Situación', 'Total USD', 'Abonado', 'Saldo', 'Último pago']
    filas = [cab]
    for p in docs:
        cobro = p.confirmada or p.es_venta
        filas.append([p.numero_visible, f'{p.fecha:%d/%m/%Y}', Paragraph(p.vendedor.nombre_visible, ESTILO_CELDA),
                      Paragraph(' · '.join(_situacion(p)), ESTILO_CELDA), usd(p.neto_usd),
                      usd(p.abonado_usd) if cobro else '—', usd(p.saldo_usd) if cobro else '—',
                      f'{p.fecha_pago:%d/%m/%Y}' if p.fecha_pago else '—'])
    t = totales(docs)
    filas.append([f'{t["n"]} documento(s)', '', '', 'TOTALES', usd(t['total']), usd(t['abonado']), usd(t['saldo']), ''])
    tabla = Table(filas, colWidths=[28 * mm, 20 * mm, 32 * mm, 74 * mm, 26 * mm, 24 * mm, 24 * mm, 22 * mm], repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('ALIGN', (4, 0), (6, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F3F4F6')]),
        ('LINEABOVE', (0, -1), (-1, -1), 1, color), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('TEXTCOLOR', (6, -1), (6, -1), colors.HexColor('#B02A37') if t['saldo'] else colors.black),
    ]))
    e.append(tabla)
    pie = pie_empresa(empresa)
    doc.build(e, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def excel(cliente, docs, filtro, desde=None, hasta=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Estado de cuenta'
    ws.append([f'Estado de cuenta · {cliente.nombre} · RIF {cliente.rif or "—"}'])
    ws.append([f'{FILTROS[filtro][0]} · {_rango(desde, hasta)}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    cab = ['N°', 'Fecha', 'Estado', 'Vendedor', 'Total USD', 'Abonado USD', 'Saldo USD', 'Pagada', 'Último pago',
           'Método', 'Entregada', 'Fecha entrega', 'Factura', 'N° control', 'Fecha factura', 'Devuelto USD']
    ws.append(cab)
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='053D74')
        c.alignment = Alignment(horizontal='center')
    for p in docs:
        cobro = p.confirmada or p.es_venta
        ws.append([p.numero_visible, p.fecha, p.get_estado_display(), p.vendedor.nombre_visible, float(p.neto_usd),
                   float(p.abonado_usd) if cobro else None, float(p.saldo_usd) if cobro else None,
                   ('Sí' if p.pagado else ('Parcial' if p.abonado_usd else 'No')) if cobro else '',
                   p.fecha_pago, p.get_metodo_pago_display() if p.metodo_pago else '',
                   ('Sí' if p.entregado else 'No') if cobro else '', p.fecha_entrega,
                   p.numero_factura or '', p.numero_control or '', p.fecha_facturacion,
                   float(p.devuelto_usd) or None])
    t = totales(docs)
    ws.append(['TOTALES', None, f'{t["n"]} documento(s)', None, float(t['total']), float(t['abonado']), float(t['saldo'])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for fila in ws.iter_rows(min_row=5):
        for c in fila:
            if c.column in (5, 6, 7, 16):
                c.number_format = '#,##0.00'
            elif c.column in (2, 9, 12, 15):
                c.number_format = 'DD/MM/YYYY'
    for letra, ancho in zip('ABCDEFGHIJKLMNOP', (16, 11, 14, 18, 12, 12, 12, 9, 12, 14, 10, 12, 12, 12, 12, 12)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'A5'
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
