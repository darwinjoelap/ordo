"""
Importación masiva de clientes desde Excel (.xlsx). Mismo flujo que la de productos:
plantilla → validar() sin guardar nada → aplicar() en una sola transacción.

Un cliente del archivo es «el mismo» que uno de Ordo si coincide el RIF; sin RIF, si coincide el nombre
(sin distinguir mayúsculas, acentos ni signos). En ese caso se actualiza en vez de duplicarse.
"""
import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.inventario.importacion import AZUL, VERDE, Columna, ErrorArchivo, _si_no, _texto

from .models import Cliente

MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024
HOJA = 'Clientes'

COLUMNAS = [
    Columna('nombre', 'nombre', True, 'Texto, máx. 200', 'Agrícola Las Raíces, C.A.', 'Nombre o razón social.', 38),
    Columna('rif', 'rif', False, 'Texto, máx. 20', 'J-12345678-9', 'RIF o cédula. Sale en los presupuestos. No se puede repetir.', 16),
    Columna('contacto', 'contacto', False, 'Texto, máx. 100', 'Carolina Pérez', 'Persona de contacto.', 22),
    Columna('telefono', 'telefono', False, 'Texto, máx. 50', '0414-5551234', 'Uno o varios, separados por « / ».', 22),
    Columna('email', 'email', False, 'Correo', 'compras@ejemplo.com', 'Opcional.', 26),
    Columna('direccion', 'direccion', False, 'Texto', 'Av. Los Pioneros, C.C. La Espiga. Acarigua, Portuguesa', 'Dirección fiscal o de entrega.', 48),
    Columna('notas', 'notas', False, 'Texto', 'Paga a 30 días', 'Notas internas: no salen en los documentos.', 28),
    Columna('vendedor', 'vendedor', False, 'Usuario de Ordo', 'maria', 'Nombre de usuario del vendedor asignado. Vacío = sin asignar.', 14),
    Columna('activo', 'activo', False, 'SI / NO', 'SI', 'Vacío = SI.', 8),
]
CLAVES = [c.clave for c in COLUMNAS]
LARGOS = {'nombre': 200, 'rif': 20, 'contacto': 100, 'telefono': 50}

REGLAS = [
    'La primera fila de la hoja "Clientes" son los títulos: no los cambies.',
    'Una fila por cliente. Solo el nombre es obligatorio.',
    'Si el RIF (o, sin RIF, el nombre) ya existe en Ordo, se actualizan sus datos en vez de crear otro cliente. '
    'Una celda vacía NO borra lo que ya tenía.',
    'Dentro del archivo no puede repetirse un RIF ni un nombre.',
    'Las filas totalmente vacías se ignoran. Máximo 5.000 filas y 5 MB por archivo.',
    'Antes de guardar nada, Ordo valida todo el archivo y te muestra los errores por fila. '
    'Si hay un solo error, no se importa nada: corrige y vuelve a subirlo.',
]


def clave_nombre(nombre):
    """«AGRÍCOLA  Las Raíces, C.A.» → «agricola las raices ca»: para reconocer el mismo cliente."""
    t = unicodedata.normalize('NFKD', str(nombre or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(re.sub(r'[^a-z0-9 ]', '', t.replace('\xa0', ' ')).split())


def clave_rif(rif):
    return re.sub(r'[^A-Z0-9]', '', str(rif or '').upper())


def generar_plantilla(vendedores=None, con_ejemplos=True, filas=None):
    """Excel con hojas Clientes e Instrucciones (y Vendedores). `filas`: lista de dicts para exportar."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    for i, col in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i, value=col.titulo)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL if col.obligatoria else VERDE)
        celda.alignment = Alignment(horizontal='center')
        celda.comment = Comment(f'{"OBLIGATORIA" if col.obligatoria else "Opcional"} · {col.formato}\n{col.nota}', 'Ordo')
        ws.column_dimensions[get_column_letter(i)].width = col.ancho
        if col.formato == 'SI / NO':
            dv = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}{MAX_FILAS + 1}')
            ws.add_data_validation(dv)
    ws.freeze_panes = 'B2'
    if filas is None and con_ejemplos:
        filas = [{c.clave: c.ejemplo for c in COLUMNAS},
                 {'nombre': 'Juan Francisco Miquelena', 'rif': 'V-12345678', 'telefono': '0424-5551234'}]
    for r, fila in enumerate(filas or [], start=2):
        for i, clave in enumerate(CLAVES, start=1):
            ws.cell(row=r, column=i, value=fila.get(clave) or None)

    wi = wb.create_sheet('Instrucciones')
    wi['A1'] = 'Cómo llenar la plantilla de clientes de Ordo'
    wi['A1'].font = Font(bold=True, size=14, color=AZUL)
    fila = 3
    for regla in REGLAS:
        wi.cell(row=fila, column=1, value=f'• {regla}').alignment = Alignment(wrap_text=True, vertical='top')
        wi.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
        wi.row_dimensions[fila].height = 32
        fila += 1
    fila += 1
    for i, t in enumerate(['Columna', '¿Obligatoria?', 'Formato', 'Ejemplo', 'Notas'], start=1):
        c = wi.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
    for col in COLUMNAS:
        fila += 1
        for i, v in enumerate([col.titulo, 'SÍ' if col.obligatoria else 'no', col.formato, col.ejemplo, col.nota], start=1):
            wi.cell(row=fila, column=i, value=v).alignment = Alignment(wrap_text=True, vertical='top')
    for letra, ancho in zip('ABCDE', (16, 13, 20, 34, 60)):
        wi.column_dimensions[letra].width = ancho

    if vendedores:
        wv = wb.create_sheet('Vendedores')
        wv.append(['usuario', 'nombre'])
        for c in wv[1]:
            c.font = Font(bold=True)
        for usuario, nombre in vendedores:
            wv.append([usuario, nombre])
        wv.column_dimensions['A'].width = 18
        wv.column_dimensions['B'].width = 30
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def exportar_clientes(clientes):
    filas = [{'nombre': c.nombre, 'rif': c.rif, 'contacto': c.contacto, 'telefono': c.telefono, 'email': c.email,
              'direccion': c.direccion, 'notas': c.notas, 'vendedor': c.vendedor.username if c.vendedor_id else '',
              'activo': 'SI' if c.activo else 'NO'} for c in clientes]
    return generar_plantilla(filas=filas)


@dataclass
class ClienteImportado:
    fila: int
    datos: dict
    existente: Cliente | None = None


@dataclass
class Resultado:
    clientes: list
    errores: list          # [(fila, columna, mensaje)]
    avisos: list           # [(fila, mensaje)]

    @property
    def valido(self):
        return not self.errores

    def resumen(self):
        nuevos = sum(1 for c in self.clientes if not c.existente)
        return {'nuevos': nuevos, 'actualizados': len(self.clientes) - nuevos,
                'sin_rif': sum(1 for c in self.clientes if not c.datos['rif'] and not (c.existente and c.existente.rif)),
                'errores': len(self.errores), 'avisos': len(self.avisos)}


def leer_filas(contenido):
    if len(contenido) > MAX_BYTES:
        raise ErrorArchivo('El archivo supera 5 MB.')
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorArchivo('No se pudo leer el archivo. Debe ser un Excel .xlsx (usa la plantilla de Ordo).')
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    try:
        encabezado = [_texto(c).lower() for c in next(filas)]
    except StopIteration:
        raise ErrorArchivo('La hoja está vacía.')
    if 'nombre' not in encabezado:
        raise ErrorArchivo('Falta la columna obligatoria «nombre». Usa la plantilla de Ordo.')
    indices = {clave: encabezado.index(clave) for clave in CLAVES if clave in encabezado}
    salida = []
    for numero, valores in enumerate(filas, start=2):
        if valores is None or all(v in (None, '') for v in valores):
            continue
        salida.append((numero, {clave: (valores[i] if i < len(valores) else None) for clave, i in indices.items()}))
        if len(salida) > MAX_FILAS:
            raise ErrorArchivo(f'El archivo tiene más de {MAX_FILAS} filas con datos.')
    wb.close()
    if not salida:
        raise ErrorArchivo('El archivo no tiene filas con datos.')
    return salida


def validar(contenido, empresa):
    """Valida el archivo completo sin guardar nada."""
    filas = leer_filas(contenido)
    errores, avisos, clientes = [], [], []
    por_rif, por_nombre = {}, {}
    for c in Cliente.objects.all():
        if c.rif:
            por_rif[clave_rif(c.rif)] = c
        por_nombre.setdefault(clave_nombre(c.nombre), c)
    vendedores = {u.username.lower(): u for u in get_user_model().objects.filter(
        membresias__empresa=empresa, membresias__activa=True)}
    rifs_vistos, nombres_vistos = {}, {}

    for numero, crudo in filas:
        d = {clave: ' '.join(_texto(crudo.get(clave)).replace('\xa0', ' ').split()) if clave not in ('direccion', 'notas')
             else _texto(crudo.get(clave)).replace('\xa0', ' ') for clave in CLAVES if clave not in ('vendedor', 'activo')}
        d['rif'] = d['rif'].upper()
        if not d['nombre']:
            errores.append((numero, 'nombre', 'Falta el nombre.'))
            continue
        for clave, maximo in LARGOS.items():
            if len(d[clave]) > maximo:
                errores.append((numero, clave, f'Máximo {maximo} caracteres (tiene {len(d[clave])}).'))
        if d['email']:
            d['email'] = d['email'].lower()
            try:
                validate_email(d['email'])
            except ValidationError:
                errores.append((numero, 'email', f'«{d["email"]}» no es un correo válido.'))
        try:
            d['activo'] = _si_no(crudo.get('activo'), None)
        except ValueError as e:
            errores.append((numero, 'activo', str(e)))
            d['activo'] = None
        d['vendedor'] = None
        usuario = _texto(crudo.get('vendedor')).lower()
        if usuario:
            d['vendedor'] = vendedores.get(usuario)
            if not d['vendedor']:
                errores.append((numero, 'vendedor', f'No hay un usuario «{usuario}» activo en la empresa.'))

        kr, kn = clave_rif(d['rif']), clave_nombre(d['nombre'])
        if kr and kr in rifs_vistos:
            errores.append((numero, 'rif', f'RIF repetido: ya está en la fila {rifs_vistos[kr]}.'))
        elif kn in nombres_vistos:
            errores.append((numero, 'nombre', f'Cliente repetido: ya está en la fila {nombres_vistos[kn]}.'))
        if kr:
            rifs_vistos[kr] = numero
        nombres_vistos.setdefault(kn, numero)

        existente = por_rif.get(kr) if kr else None
        if not existente:
            candidato = por_nombre.get(kn)
            if candidato and kr and candidato.rif and clave_rif(candidato.rif) != kr:
                avisos.append((numero, f'Ya existe «{candidato.nombre}» con otro RIF ({candidato.rif}): se creará como cliente aparte.'))
            elif candidato:
                existente = candidato
        if not kr and not (existente and existente.rif):
            avisos.append((numero, 'Sin RIF: sus documentos saldrán sin ese dato.'))
        clientes.append(ClienteImportado(numero, d, existente))
    return Resultado(clientes, errores, avisos)


@transaction.atomic
def aplicar(resultado):
    """Crea y actualiza. En los existentes, una celda vacía conserva el dato que ya había."""
    creados = actualizados = 0
    for c in resultado.clientes:
        d = c.datos
        if c.existente:
            cliente = c.existente
            for campo in ('nombre', 'rif', 'contacto', 'telefono', 'email', 'direccion', 'notas', 'vendedor'):
                if d[campo]:
                    setattr(cliente, campo, d[campo])
            if d['activo'] is not None:
                cliente.activo = d['activo']
            cliente.save()
            actualizados += 1
        else:
            Cliente.objects.create(nombre=d['nombre'], rif=d['rif'], contacto=d['contacto'], telefono=d['telefono'],
                                   email=d['email'], direccion=d['direccion'], notas=d['notas'],
                                   vendedor=d['vendedor'], activo=True if d['activo'] is None else d['activo'])
            creados += 1
    return {'creados': creados, 'actualizados': actualizados}
