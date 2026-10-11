"""
Cobranza: el dinero que entró (abonos), por fecha de cada pago. Es distinto de «ventas», que cuenta la venta
completa el día que se valida. Una venta de un mes puede cobrarse en el siguiente: su abono cuenta en el mes
en que entró el dinero.
"""
import calendar
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Q, Sum
from django.utils import timezone

from .models import Abono, Presupuesto

CERO = Decimal('0')
E = Presupuesto.Estado


def abonos_de(docs):
    """Abonos de los documentos que el usuario puede ver (docs: queryset de Presupuesto)."""
    return Abono.objects.filter(presupuesto__in=docs)


def resumen(abonos, desde, hasta):
    """Totales del período: USD equivalente, divisas, bolívares, por método y por día."""
    qs = abonos.filter(fecha__gte=desde, fecha__lte=hasta)
    filas = list(qs.values('fecha', 'moneda', 'metodo', 'monto', 'monto_usd'))
    tot = {'usd': CERO, 'divisas': CERO, 'bs': CERO, 'bs_usd': CERO, 'n': len(filas)}
    metodos, dias = defaultdict(lambda: CERO), defaultdict(lambda: CERO)
    for f in filas:
        tot['usd'] += f['monto_usd']
        if f['moneda'] == 'BS':
            tot['bs'] += f['monto']
            tot['bs_usd'] += f['monto_usd']
        else:
            tot['divisas'] += f['monto_usd']
        metodos[f['metodo']] += f['monto_usd']
        dias[f['fecha']] += f['monto_usd']
    etiquetas = dict(Presupuesto.MetodoPago.choices)
    tot['metodos'] = sorted(({'metodo': etiquetas.get(k, k), 'clave': k, 'usd': v} for k, v in metodos.items()),
                            key=lambda x: -x['usd'])
    tot['dias'] = dias
    return tot


def tablero(docs, hoy=None):
    """Cifras de la tarjeta «Dinero ingresado» del inicio."""
    hoy = hoy or timezone.localdate()
    abonos = abonos_de(docs)
    inicio = hoy.replace(day=1)
    mes = resumen(abonos, inicio, hoy)
    anterior_ini = (inicio - timedelta(days=1)).replace(day=1)
    ultimo_ant = calendar.monthrange(anterior_ini.year, anterior_ini.month)[1]
    anterior_fin = anterior_ini.replace(day=min(hoy.day, ultimo_ant))
    previo = abonos.filter(fecha__gte=anterior_ini, fecha__lte=anterior_fin).aggregate(t=Sum('monto_usd'))['t'] or CERO
    maximo = max(mes['dias'].values(), default=CERO)
    barras = [{'dia': d, 'usd': mes['dias'].get(inicio.replace(day=d), CERO),
               'alto': int(mes['dias'].get(inicio.replace(day=d), CERO) * 100 / maximo) if maximo else 0,
               'hoy': d == hoy.day}
              for d in range(1, calendar.monthrange(hoy.year, hoy.month)[1] + 1)]
    variacion = None
    if previo > 0:
        variacion = ((mes['usd'] - previo) * 100 / previo).quantize(Decimal('1'))
    return {**mes, 'hoy': mes['dias'].get(hoy, CERO), 'previo': previo, 'variacion': variacion,
            'mes_anterior': anterior_ini, 'barras': barras, 'top': mes['metodos'][:3]}


def por_cobrar(docs, hoy=None):
    """Saldo pendiente (total − abonado) y cuánto lleva más de 30 días desde que se confirmó."""
    hoy = hoy or timezone.localdate()
    pendientes = docs.filter(Q(estado=E.POR_PAGAR) | Q(estado=E.VALIDADA, pagado=False))
    limite = timezone.make_aware(timezone.datetime.combine(hoy - timedelta(days=30), timezone.datetime.min.time()))
    t = {'n': 0, 'usd': CERO, 'viejo_usd': CERO, 'viejo_n': 0}
    for p in pendientes.only('total_usd', 'devuelto_usd', 'abonado_usd', 'confirmado_en', 'validado_en'):
        saldo = p.saldo_usd
        t['n'] += 1
        t['usd'] += saldo
        desde = p.confirmado_en or p.validado_en
        if desde and desde < limite:
            t['viejo_usd'] += saldo
            t['viejo_n'] += 1
    return t


# ── Reporte ───────────────────────────────────────────────────────────────────

def detalle(abonos, desde, hasta, metodo=''):
    qs = abonos.filter(fecha__gte=desde, fecha__lte=hasta)
    if metodo:
        qs = qs.filter(metodo=metodo)
    return qs.select_related('presupuesto__cliente', 'presupuesto__vendedor', 'registrado_por').order_by('fecha', 'pk')


def pdf(filas, tot, desde, hasta, empresa, filtro=''):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Spacer, Table, TableStyle

    from apps.core.pdf import ESTILO_CELDA, Paragraph, documento, encabezado_empresa, pie_empresa
    from apps.core.templatetags.ordo import usd
    from .pdf import bs
    perfil = empresa.perfil
    color = colors.HexColor(perfil.color_principal or '#053D74')
    buffer = BytesIO()
    doc = documento(buffer, titulo_pdf='Cobranza', horizontal=True)
    e = encabezado_empresa(empresa, titulo='COBRANZA',
                           subtitulo=f'Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}' + (f'<br/>{filtro}' if filtro else ''))
    resumen_t = [['Cobrado (USD equivalente)', usd(tot['usd'])], ['En divisas', usd(tot['divisas'])],
                 ['En bolívares', f'{bs(tot["bs"])} (≈ {usd(tot["bs_usd"])})']]
    resumen_t += [[f'· {m["metodo"]}', usd(m['usd'])] for m in tot['metodos']]
    tr = Table(resumen_t, colWidths=[70 * mm, 60 * mm], hAlign='LEFT')
    tr.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 9), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                            ('ALIGN', (1, 0), (1, -1), 'RIGHT'), ('LINEBELOW', (0, 2), (-1, 2), 0.5, colors.HexColor('#D1D5DB'))]))
    e += [tr, Spacer(1, 5 * mm)]
    cab = ['Fecha', 'Venta', 'Cliente', 'Vendedor', 'Método', 'Banco / Ref.', 'Recibido', 'USD']
    datos = [cab]
    for a in filas:
        p = a.presupuesto
        recibido = bs(a.monto) + (f' @ {a.tasa_bs:.2f}'.replace('.', ',') if a.tasa_bs else '') if a.moneda == 'BS' else usd(a.monto)
        datos.append([f'{a.fecha:%d/%m/%Y}', p.numero, Paragraph(p.cliente.nombre, ESTILO_CELDA),
                      Paragraph(p.vendedor.nombre_visible, ESTILO_CELDA), a.get_metodo_display(),
                      Paragraph(' · '.join(x for x in [a.banco, a.referencia] if x) or '—', ESTILO_CELDA),
                      recibido, usd(a.monto_usd)])
    datos.append([f'{len(filas)} pago(s)', '', '', '', '', '', 'TOTAL', usd(tot['usd'])])
    t = Table(datos, colWidths=[20 * mm, 26 * mm, 58 * mm, 30 * mm, 26 * mm, 34 * mm, 34 * mm, 22 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('ALIGN', (6, 0), (7, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F3F4F6')]),
        ('LINEABOVE', (0, -1), (-1, -1), 1, color), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    e.append(t)
    pie = pie_empresa(empresa)
    doc.build(e, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def excel(filas, tot, desde, hasta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Cobranza'
    ws.append([f'Cobranza del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Cobrado USD equiv.', float(tot['usd']), 'Divisas', float(tot['divisas']), 'Bolívares', float(tot['bs'])])
    ws.append([])
    ws.append(['Fecha', 'Venta', 'Cliente', 'RIF', 'Vendedor', 'Método', 'Banco', 'Referencia', 'Moneda', 'Monto recibido',
               'Tasa', 'Equivale USD', 'Registró'])
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='053D74')
    for a in filas:
        p = a.presupuesto
        ws.append([a.fecha, p.numero, p.cliente.nombre, p.cliente.rif, p.vendedor.nombre_visible, a.get_metodo_display(),
                   a.banco, a.referencia, a.get_moneda_display(), float(a.monto), float(a.tasa_bs) if a.tasa_bs else None,
                   float(a.monto_usd), a.registrado_por.nombre_visible if a.registrado_por_id else ''])
    ws.append(['TOTAL', None, f'{len(filas)} pago(s)'] + [None] * 8 + [float(tot['usd'])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for fila in ws.iter_rows(min_row=5):
        fila[0].number_format = 'DD/MM/YYYY'
        for i in (9, 11):
            fila[i].number_format = '#,##0.00'
    for letra, ancho in zip('ABCDEFGHIJKLM', (11, 15, 30, 14, 16, 15, 16, 14, 8, 15, 10, 13, 16)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'A5'
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
