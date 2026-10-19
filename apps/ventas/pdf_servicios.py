"""PDF y Excel de servicios: constancia de servicio (con firma) y reporte de servicios realizados."""
from io import BytesIO

from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Spacer, Table, TableStyle

from apps.core.pdf import ESTILO_CELDA, Paragraph, documento, encabezado_empresa, pie_empresa
from apps.core.templatetags.ordo import usd

from .pdf import E_CENTRO, E_NORMAL, E_PEQUENO, GRIS_BORDE, GRIS_CLARO, _estilo, _imagen_firma


def constancia_pdf(item, empresa):
    """Constancia de servicio: qué se hizo, cuándo, quién lo hizo y quién lo recibió en el cliente."""
    p = item.presupuesto
    perfil = empresa.perfil
    color = colors.HexColor(perfil.color_principal or '#053D74')
    buffer = BytesIO()
    numero = f'{p.numero}-S{item.pk}'
    doc = documento(buffer, titulo_pdf=f'Constancia de servicio {numero}')
    cuando = timezone.localtime(item.realizado_en) if item.realizado_en else None
    sub = [f'<font size="12"><b>N° {numero}</b></font>', '', f'Venta: {p.numero}']
    sub.append(f'Realizado: {cuando:%d/%m/%Y %H:%M}' if cuando else
               (f'Programado: {item.fecha_servicio:%d/%m/%Y}' if item.fecha_servicio else 'Por programar'))
    e = encabezado_empresa(empresa, titulo='CONSTANCIA DE SERVICIO', subtitulo='<br/>'.join(sub))
    c = p.cliente
    e += [Paragraph(f'<b>Cliente:</b> {c.nombre} · RIF/Cédula: {c.rif or "—"}', E_NORMAL)]
    if c.direccion:
        e.append(Paragraph(f'<b>Dirección:</b> {c.direccion}', E_NORMAL))
    e += [Paragraph(f'<b>Responsable:</b> {item.responsable_visible or "—"}', E_NORMAL), Spacer(1, 4 * mm)]

    hi = _estilo('s_hi', textColor=colors.white, fontName='Helvetica-Bold')
    hc = _estilo('s_hc', textColor=colors.white, fontName='Helvetica-Bold', alignment=TA_CENTER)
    desc = item.producto.nombre
    if item.descripcion_tecnica:
        desc += '<br/><font size="7.5" color="#495057">' + item.descripcion_tecnica.strip().replace('\n', '<br/>') + '</font>'
    filas = [[Paragraph('Cód.', hi), Paragraph('Servicio', hi), Paragraph('Cantidad', hc), Paragraph('Und.', hc)],
             [Paragraph(item.producto.codigo, E_NORMAL), Paragraph(desc, E_NORMAL),
              Paragraph(str(item.cantidad), E_CENTRO), Paragraph(item.producto.unidad.abreviatura, E_CENTRO)]]
    t = Table(filas, colWidths=[28 * mm, 112 * mm, 24 * mm, 22 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, GRIS_CLARO]),
        ('BOX', (0, 0), (-1, -1), 0.5, GRIS_BORDE), ('INNERGRID', (0, 0), (-1, -1), 0.25, GRIS_BORDE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    e += [t, Spacer(1, 4 * mm)]
    if item.nota_servicio:
        e += [Paragraph(f'<b>Observaciones de la ejecución:</b> {item.nota_servicio}', E_NORMAL), Spacer(1, 4 * mm)]

    fecha = f'{cuando:%d/%m/%Y %H:%M}' if cuando else 'Fecha: ____ / ____ / ________'

    def bloque(titulo, nombre, cedula, firma):
        partes = [Paragraph(f'<b>{titulo}</b>', E_NORMAL)]
        img = _imagen_firma(firma)
        partes.append(img if img else Paragraph('<br/><br/>______________________________', E_NORMAL))
        partes.append(Paragraph(f'{nombre or "Nombre: ____________________"}<br/>C.I.: {cedula or "__________"}<br/>{fecha}',
                                E_PEQUENO))
        return partes
    hizo = item.realizado_por.nombre_visible if item.realizado_por_id else item.responsable_visible
    celdas = [[bloque('Realizó', hizo, '', ''),
               bloque('Conforme el cliente', item.recibe_nombre, item.recibe_cedula, item.recibe_firma)]]
    firmas = Table(celdas, colWidths=[93 * mm, 93 * mm])
    firmas.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOX', (0, 0), (-1, -1), 0.5, GRIS_BORDE),
                                ('INNERGRID', (0, 0), (-1, -1), 0.5, GRIS_BORDE), ('LEFTPADDING', (0, 0), (-1, -1), 6)]))
    e.append(KeepTogether([firmas]))
    pie = pie_empresa(empresa)
    doc.build(e, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def reporte_pdf(datos, desde, hasta, empresa):
    perfil = empresa.perfil
    color = colors.HexColor(perfil.color_principal or '#053D74')
    buffer = BytesIO()
    doc = documento(buffer, titulo_pdf='Servicios realizados', horizontal=True)
    e = encabezado_empresa(empresa, titulo='SERVICIOS REALIZADOS', subtitulo=f'Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}')
    resumen = [['Servicios realizados', str(datos['n'])], ['Monto (USD, antes de descuento e IVA)', usd(datos['total'])],
               ['Cobrado', usd(datos['cobrado'])], ['Por cobrar', usd(datos['por_cobrar'])]]
    tr = Table(resumen, colWidths=[75 * mm, 40 * mm], hAlign='LEFT')
    tr.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 9), ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold')]))

    def tabla_resumen(titulo, filas, con_cantidad):
        cab = [titulo, 'N°'] + (['Cant.'] if con_cantidad else []) + ['USD']
        datos_t = [cab] + [[Paragraph(f['nombre'], ESTILO_CELDA), str(f['n'])]
                           + ([str(f['cantidad'])] if con_cantidad else []) + [usd(f['monto'])] for f in filas]
        anchos = [70 * mm, 12 * mm] + ([16 * mm] if con_cantidad else []) + [24 * mm]
        t = Table(datos_t, colWidths=anchos, hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), color), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                               ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
                               ('ALIGN', (1, 0), (-1, -1), 'RIGHT'), ('GRID', (0, 0), (-1, -1), 0.25, GRIS_BORDE)]))
        return t
    lado = Table([[tr, tabla_resumen('Servicio', datos['por_servicio'], True),
                   tabla_resumen('Responsable', datos['por_responsable'], False)]],
                 colWidths=[90 * mm, 95 * mm, 80 * mm])
    lado.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    e += [lado, Spacer(1, 5 * mm)]

    cab = ['Fecha', 'Venta', 'Cliente', 'Servicio', 'Cant.', 'Responsable', 'USD', 'Cobro']
    filas = [cab]
    for i in datos['filas']:
        p = i.presupuesto
        filas.append([f'{timezone.localtime(i.realizado_en):%d/%m/%Y}', p.numero, Paragraph(p.cliente.nombre, ESTILO_CELDA),
                      Paragraph(i.producto.nombre, ESTILO_CELDA), str(i.cantidad),
                      Paragraph(i.responsable_visible or '—', ESTILO_CELDA), usd(i.monto),
                      'Cobrado' if i.cobrada else 'Por cobrar'])
    filas.append([f'{datos["n"]} servicio(s)', '', '', '', '', 'TOTAL', usd(datos['total']), ''])
    t = Table(filas, colWidths=[20 * mm, 26 * mm, 58 * mm, 60 * mm, 13 * mm, 38 * mm, 22 * mm, 20 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('ALIGN', (4, 0), (4, -1), 'CENTER'), ('ALIGN', (6, 0), (6, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F3F4F6')]),
        ('LINEABOVE', (0, -1), (-1, -1), 1, color), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    e.append(t)
    pie = pie_empresa(empresa)
    doc.build(e, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def reporte_excel(datos, desde, hasta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Servicios'
    ws.append([f'Servicios realizados del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Realizados', datos['n'], 'Monto USD', float(datos['total']), 'Cobrado', float(datos['cobrado']),
               'Por cobrar', float(datos['por_cobrar'])])
    ws.append([])
    ws.append(['Fecha', 'Venta', 'Cliente', 'RIF', 'Código', 'Servicio', 'Cantidad', 'Unidad', 'Precio USD', 'Monto USD',
               'Responsable', 'Realizó', 'Recibió', 'Nota', 'Cobro'])
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='053D74')
    for i in datos['filas']:
        p = i.presupuesto
        ws.append([timezone.localtime(i.realizado_en).date(), p.numero, p.cliente.nombre, p.cliente.rif,
                   i.producto.codigo, i.producto.nombre, i.cantidad, i.producto.unidad.abreviatura, float(i.precio_usd),
                   float(i.monto), i.responsable_visible, i.realizado_por.nombre_visible if i.realizado_por_id else '',
                   i.recibe_nombre, i.nota_servicio, 'Cobrado' if i.cobrada else 'Por cobrar'])
    ws.append(['TOTAL', None, f'{datos["n"]} servicio(s)'] + [None] * 6 + [float(datos['total'])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for fila in ws.iter_rows(min_row=5):
        fila[0].number_format = 'DD/MM/YYYY'
        for n in (8, 9):
            fila[n].number_format = '#,##0.00'
    for letra, ancho in zip('ABCDEFGHIJKLMNO', (11, 15, 30, 14, 12, 30, 9, 8, 11, 12, 20, 18, 20, 30, 11)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'A5'

    for titulo, filas, cols in (('Por servicio', datos['por_servicio'], ('Servicio', 'N°', 'Cantidad', 'Monto USD')),
                                ('Por responsable', datos['por_responsable'], ('Responsable', 'N°', 'Monto USD'))):
        h = wb.create_sheet(titulo)
        h.append(list(cols))
        for c in h[1]:
            c.font = Font(bold=True)
        for f in filas:
            h.append([f['nombre'], f['n']] + ([f['cantidad']] if 'Cantidad' in cols else []) + [float(f['monto'])])
        h.column_dimensions['A'].width = 34
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
